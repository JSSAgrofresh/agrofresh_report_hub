"""Utilidades para armar el documento Word con la marca AgroFresh.

Solo envuelve python-docx: estilos, portada con fondo, encabezado con logo, pie con
número de página, títulos numerados, tablas con cabecera de marca, figuras con leyenda y
páginas horizontales para los diagramas anchos.
"""
from __future__ import annotations

import copy
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Pt, RGBColor

VERDE = RGBColor(0x43, 0xB0, 0x49)
VERDE_OSCURO = RGBColor(0x2F, 0x5F, 0x3A)
AMARILLO = RGBColor(0xF0, 0xC4, 0x19)
GRIS = RGBColor(0x55, 0x5F, 0x58)
FUENTE = "Calibri"

HEX_VERDE = "43B049"
HEX_VERDE_CLARO = "E9F5E7"
HEX_AMARILLO_CLARO = "FDF3CF"
HEX_GRIS_CLARO = "F4F6F4"


def _fuente(estilo, tam, color=None, negrita=None, cursiva=None):
    estilo.font.name = FUENTE
    estilo.element.rPr.rFonts.set(qn("w:eastAsia"), FUENTE)
    estilo.font.size = Pt(tam)
    if color is not None:
        estilo.font.color.rgb = color
    if negrita is not None:
        estilo.font.bold = negrita
    if cursiva is not None:
        estilo.font.italic = cursiva


def nuevo_documento() -> Document:
    doc = Document()
    est = doc.styles
    _fuente(est["Normal"], 11)
    est["Normal"].paragraph_format.space_after = Pt(6)
    est["Normal"].paragraph_format.line_spacing = 1.2
    for nombre, tam, color, antes, despues in (
        ("Heading 1", 22, VERDE_OSCURO, 0, 14),
        ("Heading 2", 15, VERDE_OSCURO, 16, 6),
        ("Heading 3", 12.5, VERDE, 12, 4),
    ):
        _fuente(est[nombre], tam, color, negrita=True)
        est[nombre].paragraph_format.space_before = Pt(antes)
        est[nombre].paragraph_format.space_after = Pt(despues)
        est[nombre].paragraph_format.keep_with_next = True
    _fuente(est["Caption"], 9.5, GRIS, negrita=False, cursiva=True)
    est["Caption"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    est["Caption"].paragraph_format.space_after = Pt(12)
    for nombre in ("List Bullet", "List Number"):
        _fuente(est[nombre], 11)
        est[nombre].paragraph_format.space_after = Pt(3)
    # Estilo para el código / comandos
    cod = est.add_style("Codigo", 1)
    cod.base_style = est["Normal"]
    cod.font.name = "Consolas"
    cod.element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
    cod.font.size = Pt(9)
    cod.paragraph_format.left_indent = Cm(0.5)
    cod.paragraph_format.space_after = Pt(0)
    cod.paragraph_format.line_spacing = 1.0
    # Márgenes
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21.0), Cm(29.7)
    s.left_margin = s.right_margin = Cm(2.3)
    s.top_margin, s.bottom_margin = Cm(2.4), Cm(2.2)
    return doc


# ---------------------------------------------------------------------------
# Campos de Word (número de página, tabla de contenido)
# ---------------------------------------------------------------------------

def _fld_char(run, tipo):
    el = OxmlElement("w:fldChar")
    el.set(qn("w:fldCharType"), tipo)
    run._r.append(el)


def agregar_campo(parrafo, instruccion, texto_cache="1", negrita=False, tam=None, color=None):
    r1 = parrafo.add_run()
    _fld_char(r1, "begin")
    r2 = parrafo.add_run()
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = f" {instruccion} "
    r2._r.append(it)
    r3 = parrafo.add_run()
    _fld_char(r3, "separate")
    r4 = parrafo.add_run(texto_cache)
    r5 = parrafo.add_run()
    _fld_char(r5, "end")
    for r in (r1, r2, r3, r4, r5):
        r.font.bold = negrita
        if tam:
            r.font.size = Pt(tam)
        if color is not None:
            r.font.color.rgb = color
    return r4


# ---------------------------------------------------------------------------
# Portada con la imagen de marca de fondo
# ---------------------------------------------------------------------------

def _ancla_de_pagina(run, ancho_cm, alto_cm):
    """Convierte la imagen en línea de `run` en una imagen detrás del texto, pegada a la página."""
    inline = run._r.xpath(".//wp:inline")[0]
    extent = inline.find(qn("wp:extent"))
    cx, cy = int(Cm(ancho_cm)), int(Cm(alto_cm))
    extent.set("cx", str(cx))
    extent.set("cy", str(cy))
    xfrm_ext = inline.xpath(".//a:ext")
    for e in xfrm_ext:
        if e.get("cx"):
            e.set("cx", str(cx))
            e.set("cy", str(cy))
    anchor = parse_xml(
        '<wp:anchor %s distT="0" distB="0" distL="0" distR="0" simplePos="0" relativeHeight="0" '
        'behindDoc="1" locked="0" layoutInCell="1" allowOverlap="1">'
        '<wp:simplePos x="0" y="0"/>'
        '<wp:positionH relativeFrom="page"><wp:posOffset>0</wp:posOffset></wp:positionH>'
        '<wp:positionV relativeFrom="page"><wp:posOffset>0</wp:posOffset></wp:positionV>'
        "<wp:extent cx=\"%d\" cy=\"%d\"/><wp:effectExtent l=\"0\" t=\"0\" r=\"0\" b=\"0\"/>"
        "<wp:wrapNone/></wp:anchor>" % (nsdecls("wp"), cx, cy)
    )
    for hijo in list(inline):
        if hijo.tag in (qn("wp:docPr"), qn("wp:cNvGraphicFramePr"), qn("a:graphic")):
            anchor.append(copy.deepcopy(hijo))
    inline.getparent().replace(inline, anchor)


def portada(doc, imagen_fondo: Path, titulo, subtitulo, linea_version, notas):
    s = doc.sections[0]
    s.orientation = WD_ORIENT.LANDSCAPE
    s.page_width, s.page_height = Cm(29.7), Cm(21.0)
    s.left_margin, s.right_margin = Cm(2.4), Cm(6.0)
    s.top_margin, s.bottom_margin = Cm(1.5), Cm(1.5)
    p = doc.paragraphs[0] if doc.paragraphs else doc.add_paragraph()
    r = p.add_run()
    r.add_picture(str(imagen_fondo), width=Cm(29.7))
    _ancla_de_pagina(r, 29.7, 29.7 * 1125 / 2000)
    p.paragraph_format.space_after = Pt(0)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(190)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("DOCUMENTACIÓN DEL SISTEMA")
    r.font.size, r.font.bold, r.font.color.rgb = Pt(13), True, VERDE
    r.font.name = FUENTE
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(titulo)
    r.font.size, r.font.bold, r.font.color.rgb = Pt(36), True, VERDE_OSCURO
    r.font.name = FUENTE
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(18)
    r = p.add_run(subtitulo)
    r.font.size, r.font.color.rgb = Pt(15), GRIS
    r.font.name = FUENTE
    for texto in notas:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(texto)
        r.font.size, r.font.color.rgb = Pt(11), GRIS
        r.font.name = FUENTE
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    r = p.add_run(linea_version)
    r.font.size, r.font.bold, r.font.color.rgb = Pt(11), True, VERDE_OSCURO
    r.font.name = FUENTE


# ---------------------------------------------------------------------------
# Secciones (orientación) con encabezado y pie
# ---------------------------------------------------------------------------

def _limpiar(parrafo):
    for r in list(parrafo.runs):
        r._r.getparent().remove(r._r)


def configurar_encabezado(seccion, logo: Path, texto_derecha: str):
    seccion.header.is_linked_to_previous = False
    seccion.footer.is_linked_to_previous = False
    ancho = seccion.page_width - seccion.left_margin - seccion.right_margin
    h = seccion.header.paragraphs[0]
    _limpiar(h)
    _tabs_propias(h, ancho)
    h.add_run().add_picture(str(logo), height=Cm(0.95))
    r = h.add_run("\t" + texto_derecha)
    r.font.size, r.font.color.rgb, r.font.name = Pt(9), GRIS, FUENTE
    _borde_inferior(h, HEX_VERDE)
    f = seccion.footer.paragraphs[0]
    _limpiar(f)
    _tabs_propias(f, ancho)
    r = f.add_run("AgroFresh Report Hub · Documentación del sistema\tPágina ")
    r.font.size, r.font.color.rgb, r.font.name = Pt(9), GRIS, FUENTE
    agregar_campo(f, "PAGE", "1", tam=9, color=GRIS)
    r = f.add_run(" de ")
    r.font.size, r.font.color.rgb, r.font.name = Pt(9), GRIS, FUENTE
    agregar_campo(f, "NUMPAGES", "1", tam=9, color=GRIS)


def _tabs_propias(parrafo, ancho):
    """Quita las tabulaciones centro/derecha del estilo y deja solo una a la derecha."""
    from docx.shared import Twips

    ts = parrafo.paragraph_format.tab_stops
    ts.add_tab_stop(Twips(4680), WD_TAB_ALIGNMENT.CLEAR)
    ts.add_tab_stop(Twips(9360), WD_TAB_ALIGNMENT.CLEAR)
    ts.add_tab_stop(ancho, WD_TAB_ALIGNMENT.RIGHT)


def _borde_inferior(parrafo, color_hex, grosor="8"):
    ppr = parrafo._p.get_or_add_pPr()
    borde = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    b.set(qn("w:val"), "single")
    b.set(qn("w:sz"), grosor)
    b.set(qn("w:space"), "4")
    b.set(qn("w:color"), color_hex)
    borde.append(b)
    ppr.append(borde)


def nueva_seccion(doc, horizontal: bool, logo: Path, texto_derecha: str):
    s = doc.add_section(WD_SECTION.NEW_PAGE)
    if horizontal:
        s.orientation = WD_ORIENT.LANDSCAPE
        s.page_width, s.page_height = Cm(29.7), Cm(21.0)
        s.left_margin = s.right_margin = Cm(1.8)
        s.top_margin, s.bottom_margin = Cm(2.2), Cm(1.8)
    else:
        s.orientation = WD_ORIENT.PORTRAIT
        s.page_width, s.page_height = Cm(21.0), Cm(29.7)
        s.left_margin = s.right_margin = Cm(2.3)
        s.top_margin, s.bottom_margin = Cm(2.4), Cm(2.2)
    configurar_encabezado(s, logo, texto_derecha)
    return s


# ---------------------------------------------------------------------------
# Texto
# ---------------------------------------------------------------------------

def titulo(doc, texto, nivel=1, salto_pagina=None):
    if salto_pagina is None:
        salto_pagina = nivel == 1
    if salto_pagina:
        doc.add_page_break()
    return doc.add_heading(texto, level=nivel)


def parrafo(doc, texto, negrita_inicial=None, cursiva=False, alineado=True):
    p = doc.add_paragraph()
    if alineado:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if negrita_inicial:
        r = p.add_run(negrita_inicial)
        r.bold = True
    r = p.add_run(texto)
    r.italic = cursiva
    return p


def viñetas(doc, items, estilo="List Bullet"):
    """Viñetas o lista numerada. La numerada se escribe a mano para que cada lista parta en 1."""
    for n, it in enumerate(items, 1):
        if estilo == "List Number":
            p = doc.add_paragraph()
            pf = p.paragraph_format
            pf.left_indent, pf.first_line_indent = Cm(0.9), Cm(-0.9)
            pf.space_after = Pt(3)
            pf.tab_stops.add_tab_stop(Cm(0.9))
            r = p.add_run(f"{n}.\t")
            r.bold = True
            r.font.color.rgb = VERDE_OSCURO
        else:
            p = doc.add_paragraph(style=estilo)
        if isinstance(it, tuple):
            r = p.add_run(it[0])
            r.bold = True
            p.add_run(it[1])
        else:
            p.add_run(it)


def bloque_codigo(doc, lineas):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    celda = t.rows[0].cells[0]
    _sombrear(celda, "F1F3F1")
    celda.paragraphs[0].style = doc.styles["Codigo"]
    primero = True
    for ln in lineas:
        p = celda.paragraphs[0] if primero else celda.add_paragraph(style="Codigo")
        primero = False
        p.add_run(ln)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def nota(doc, texto, titulo_nota="Nota", color_hex=HEX_AMARILLO_CLARO):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = t.rows[0].cells[0]
    _sombrear(c, color_hex)
    p = c.paragraphs[0]
    r = p.add_run(titulo_nota + ". ")
    r.bold = True
    p.add_run(texto)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


# ---------------------------------------------------------------------------
# Tablas y figuras
# ---------------------------------------------------------------------------

def _sombrear(celda, hex_color):
    tcpr = celda._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:color"), "auto")
    sh.set(qn("w:fill"), hex_color)
    tcpr.append(sh)


def _bordes_tabla(tabla, color="B9D9B5"):
    tbl = tabla._tbl
    tblpr = tbl.tblPr
    bordes = OxmlElement("w:tblBorders")
    for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
        b = OxmlElement(f"w:{lado}")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), "4")
        b.set(qn("w:space"), "0")
        b.set(qn("w:color"), color)
        bordes.append(b)
    tblpr.append(bordes)


class Contadores:
    figuras: list[tuple[int, str]] = []
    tablas: list[tuple[int, str]] = []


def leyenda(doc, tipo, texto):
    lista = Contadores.figuras if tipo == "Figura" else Contadores.tablas
    n = len(lista) + 1
    lista.append((n, texto))
    p = doc.add_paragraph(style="Caption")
    r = p.add_run(f"{tipo} ")
    r.bold = True
    agregar_campo(p, f"SEQ {tipo} \\* ARABIC", str(n), negrita=True)
    r = p.add_run(f". {texto}")
    return p


def tabla(doc, encabezados, filas, anchos_cm=None, titulo_tabla=None, tam=9.5, alinear=None):
    if titulo_tabla:
        # La leyenda de una tabla va arriba.
        p = leyenda(doc, "Tabla", titulo_tabla)
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_after = Pt(4)
    t = doc.add_table(rows=1, cols=len(encabezados))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    _bordes_tabla(t)
    t.autofit = False
    for i, h in enumerate(encabezados):
        c = t.rows[0].cells[i]
        _sombrear(c, HEX_VERDE)
        c.paragraphs[0].paragraph_format.space_after = Pt(2)
        r = c.paragraphs[0].add_run(h)
        r.bold, r.font.size, r.font.color.rgb = True, Pt(tam), RGBColor(255, 255, 255)
    # Repetir cabecera en cada página
    trpr = t.rows[0]._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader")
    th.set(qn("w:val"), "true")
    trpr.append(th)
    for n, fila in enumerate(filas):
        cells = t.add_row().cells
        # no partir una fila entre páginas
        trpr = t.rows[-1]._tr.get_or_add_trPr()
        cs = OxmlElement("w:cantSplit")
        cs.set(qn("w:val"), "true")
        trpr.append(cs)
        for i, v in enumerate(fila):
            if n % 2 == 1:
                _sombrear(cells[i], HEX_GRIS_CLARO)
            par = cells[i].paragraphs[0]
            par.paragraph_format.space_after = Pt(2)
            par.paragraph_format.line_spacing = 1.05
            r = par.add_run(str(v))
            r.font.size = Pt(tam)
            if alinear and alinear[i] == "c":
                par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif alinear and alinear[i] == "r":
                par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if anchos_cm:
        for i, w in enumerate(anchos_cm):
            t.columns[i].width = Cm(w)
        for fila in t.rows:
            for i, w in enumerate(anchos_cm):
                fila.cells[i].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    return t


def figura(doc, ruta: Path, texto, ancho_cm=None, alto_max_cm=None):
    from PIL import Image

    with Image.open(ruta) as im:
        w, h = im.size
    if ancho_cm is None:
        ancho_cm = 16.0
    if alto_max_cm and ancho_cm * h / w > alto_max_cm:
        ancho_cm = alto_max_cm * w / h
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(ruta), width=Cm(ancho_cm))
    leyenda(doc, "Figura", texto)


def salto_de_pagina(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
