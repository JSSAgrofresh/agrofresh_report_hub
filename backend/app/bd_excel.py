"""
Descarga de la BD de resultados (Report → Laboratorio) en Excel.

Mismo formato que la matriz de Solicitudes (`solicitud_excel.py`): dos filas de
encabezado -banda por laboratorio y nombre de columna- y UNA FILA POR
SOLICITUD. La diferencia es lo que hay en las columnas de analitos:

    Solicitudes   ✓ = se pidió ese analito           «FDL Dosis» = dosis anotada
    BD            el RESULTADO medido del analito    «FDL Dosis» = dosis aplicada

Las columnas se acotan a lo que hay en la descarga: un grupo o un analito que
no tiene ni un dato entre las filas descargadas no aparece. Así, si se filtra
por un laboratorio salen solo sus analitos, y si se filtra por un ingrediente
sale solo ese (más su dosis). Las columnas generales son siempre las mismas.

`construir_workbook_bd` es pura (sin base ni request): recibe las filas ya
armadas, para poder probarla sola. La consulta vive en `reportes.py`.
"""
from __future__ import annotations

from datetime import date, datetime
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .columnas_base import CAMPOS_FUNGICIDAS, GENERALES_BASE, a_fecha, partir_recepcion
from .solicitud_excel import (
    _BORDE_COMPLETO,
    VERDE_CLARO,
    VERDE_OSCURO,
    expandir_por_posicion,
    _grupos_exportacion,
)

# Columnas generales de la BD: son las MISMAS que las de la base «con muestra» de
# AgroFresh Lab, con el mismo nombre y orden, para poder cruzarlas (ver
# `columnas_base.py`, que es la única definición).
GENERALES_BD = GENERALES_BASE

# La ingesta guarda algunos analitos con un código y el catálogo de Toma de
# muestras usa otro para el MISMO analito (mapeo.py vs ANALITOS_DEFECTO). Sin
# esta tabla, el resultado de "E. Coli" caería en «Otros» en vez de en su
# columna de siempre.
ALIAS_CODIGO: dict[str, str] = {
    "ECOLI": "ECOLI100",
    "COLT": "COLIF100",
    "HONG": "HONGOS",
    "COLA": "COLIFG",
    "ECOLA": "ECOLIG",
    "ENTB": "ENTERO",
    "CEN": "CENIZAS",
    "AFLA": "AFLAT",
}

TITULO_FUNGICIDAS_BASE = "ANÁLISIS DE RESIDUOS DE FUNGICIDAS"
_PREFIJO_FUNGICIDAS = "QUITECA / AGROFRESH"


def titulo_fungicidas(laboratorios: set[str]) -> str:
    """Banda del grupo de fungicidas según los laboratorios que hay en la
    descarga: si solo está Quiteca, dice Quiteca; si están los dos, los dos.
    (En Solicitudes la banda es fija porque ahí el formato es el mismo para todos.)"""
    presentes = [nombre for nombre in ("QUITECA", "AGROFRESH") if nombre in laboratorios]
    rotulo = " / ".join(presentes) if presentes else _PREFIJO_FUNGICIDAS
    return f"{rotulo} — {TITULO_FUNGICIDAS_BASE}"


def _laboratorios_de(valor: Any) -> set[str]:
    """{'QUITECA', 'AGROFRESH'} que menciona el nombre de laboratorio de una fila
    («Quiteca», «AGROFRESH», «Quiteca / AgroFresh»...)."""
    texto = str(valor or "").casefold()
    return {n for n in ("QUITECA", "AGROFRESH") if n.casefold() in texto}


TITULO_OTROS = "OTROS ANALITOS (sin columna propia en el formato)"
# Estas columnas del formato de Solicitudes no tienen dato en la base o se
# reemplazan por «Otros»: no se muestran.
_CAMPOS_OMITIDOS = {
    "Analito Pesticida 1", "Resultado Pesticida 1",
    "Analito Pesticida 2", "Resultado Pesticida 2",
    "Analito Pesticida 3", "Resultado Pesticida 3",
}

Columna = tuple[str, str, str]  # (tipo, clave, encabezado)

# Solicitante es siempre AgroFresh (regla del laboratorio).
SOLICITANTE = "AGROFRESH"

# Lo que la base no trae se toma de la solicitud de Toma de muestras que la
# originó (se une por el N° de OT). Clave de la fila de la BD -> clave en la
# solicitud. La base manda: solo se rellena lo que quedó vacío.
DESDE_SOLICITUD: dict[str, str] = {
    "nro_solicitud": "numero_solicitud",
    "fecha_solicitud": "fecha_solicitud",
    "fecha_muestreo": "fecha_muestreo",
    "hora_muestreo": "hora_muestreo",
    "especie": "especie",
    "variedad": "variedad",
    "linea_proceso": "linea_proceso",
    "csg": "csg_productor",
    "lote": "lote",
    "posicion_muestreo": "posicion_muestreo",
    "numero_camara": "numero_camara",
    "kilos_procesados": "kilos_procesados",
    "producto_utilizado": "producto_utilizado",
    "tipo_muestra": "tipo_muestra",
    "nombre_muestreador": "nombre_muestreador",
    "generado_por": "generado_por",
    "email_solicitante": "email_solicitante",
    "email_laboratorio": "email_laboratorio",
    "csg_packing": "csg_packing",
    "codigo_muestra": "codigo_muestra",
    "peso_extraido": "peso_muestra_extraido",
}

_FECHAS_BD = {
    "fecha_solicitud", "fecha_muestreo", "fecha_entrada", "fecha_recepcion", "fecha_informe", "fecha_analisis",
}


def _vacio(v: Any) -> bool:
    return v is None or (isinstance(v, str) and not v.strip())


_a_fecha = a_fecha


def buscar_por_parecido(fila: dict[str, Any], candidatos: list[dict[str, Any]]) -> dict | None:
    """La solicitud de Toma de muestras que origina esta fila cuando la base no
    guardó su N° de OT: misma fecha de muestreo, laboratorio, sucursal y especie.

    Solo se acepta si hay **una única** coincidencia. Si hay dos (dos muestras de
    la misma planta el mismo día) no se sabe cuál es cuál y se devuelve None:
    una celda vacía es mejor que el muestreador equivocado.

    `candidatos`: filas del índice con numero_solicitud, laboratorio, ship_to,
    especie, fecha_muestreo y datos."""
    from .listados import clave_normalizada as clave

    fecha = fila.get("fecha_muestreo")
    if not isinstance(fecha, (date, datetime)) or not fila.get("ship_to") or not fila.get("especie"):
        return None
    if isinstance(fecha, datetime):
        fecha = fecha.date()
    iguales: dict[str, dict] = {}
    for c in candidatos:
        f = c.get("fecha_muestreo")
        if isinstance(f, datetime):
            f = f.date()
        if (
            f == fecha
            and clave(c.get("laboratorio") or "") == clave(fila.get("laboratorio") or "")
            and clave(c.get("ship_to") or "") == clave(fila["ship_to"])
            and clave(c.get("especie") or "") == clave(fila["especie"])
        ):
            iguales[str(c.get("numero_solicitud"))] = c["datos"]
    return next(iter(iguales.values())) if len(iguales) == 1 else None


def completar_fila(fila: dict[str, Any], datos: dict | None, correos_laboratorio: list[str] | None = None) -> None:
    """Llena, en su lugar, lo que la base no trae de una solicitud.

    `datos` es la solicitud de Toma de muestras que la originó (None si no se
    encontró: la fila queda con lo que tenga). `correos_laboratorio` son los
    contactos del laboratorio, que van en «Email Laboratorio» separados por «;»."""
    for clave, clave_solicitud in DESDE_SOLICITUD.items():
        if _vacio(fila.get(clave)) and datos and not _vacio(datos.get(clave_solicitud)):
            fila[clave] = datos[clave_solicitud]
    if _vacio(fila.get("email_laboratorio")) and correos_laboratorio:
        fila["email_laboratorio"] = "; ".join(correos_laboratorio)
    if datos:
        # Qué analitos se pidieron y con qué dosis (campos_laboratorio): la base no lo
        # guarda para las solicitudes que no pasaron por la ingesta con dosis.
        fila["_solicitud"] = {
            "analitos_solicitados": list(datos.get("analitos_solicitados") or []),
            "campos_laboratorio": dict(datos.get("campos_laboratorio") or {}),
        }
        # La recepción es el momento del cruce con la muestra (Ingreso al laboratorio).
        dia, hora = partir_recepcion(datos.get("recepcion_en"))
        if _vacio(fila.get("fecha_recepcion")) and dia:
            fila["fecha_recepcion"] = dia
        if _vacio(fila.get("hora_recepcion")) and hora:
            fila["hora_recepcion"] = hora
        # Gasto y datos del ensayo: la base solo trae el ensayo; el resto está en la solicitud.
        campos = datos.get("campos_laboratorio") or {}
        for clave, etiqueta in CAMPOS_FUNGICIDAS:
            if _vacio(fila.get(clave)) and not _vacio(campos.get(etiqueta)):
                fila[clave] = campos[etiqueta]
    for clave in _FECHAS_BD:
        fila[clave] = _a_fecha(fila.get(clave))
    fila["solicitante"] = SOLICITANTE
    if _vacio(fila.get("temporada")):
        # La temporada es el año de la muestra (2026), no un dato que se tipee.
        for clave in ("fecha_muestreo", "fecha_entrada", "fecha_informe", "fecha_analisis"):
            f = fila.get(clave)
            if isinstance(f, (date, datetime)):
                fila["temporada"] = f.year
                break



def completar_solicitados(filas: list[dict[str, Any]], analitos: list[dict]) -> None:
    """Pone en cada fila `solicitados` (códigos de formato que la solicitud pidió) y
    completa `dosis` con la que anotó la solicitud cuando la base no la trae.
    La dosis de la base manda; la «Solicitado» sin dosis anotada no es una dosis."""
    from .solicitud_excel import _analitos_fungicidas, _valor_guardado

    catalogo = _analitos_fungicidas(analitos)
    for fila in filas:
        sol = fila.get("_solicitud") or {}
        pedidos = {codigo_de_formato(c) for c in sol.get("analitos_solicitados") or []}
        fila["solicitados"] = pedidos
        campos = sol.get("campos_laboratorio") or {}
        if not campos:
            continue
        dosis = dict(fila.get("dosis") or {})
        ya = {codigo_de_formato(k) for k in dosis}
        for a in catalogo:
            cod = codigo_de_formato(a.get("codigo"))
            if cod in pedidos and cod not in ya:
                v = _valor_guardado(campos, a)
                if v and v != "Solicitado":
                    dosis[cod] = v
        fila["dosis"] = dosis
        if _vacio(fila.get("tipo_aplicacion")) and not _vacio(campos.get("Tipo Aplicación")):
            fila["tipo_aplicacion"] = campos["Tipo Aplicación"]


def codigo_de_formato(codigo: str | None) -> str:
    """El código con el que el formato de Solicitudes conoce a este analito."""
    c = (codigo or "").strip()
    return ALIAS_CODIGO.get(c.upper(), c)


def _valor_celda(v: Any) -> Any:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, (date, int, float, str)) or v is None:
        return v if v != "" else None
    return str(v)  # Decimal y demás: openpyxl no los escribe


def _numero(v: Any) -> Any:
    """Decimal/str numérico -> número; lo que no lo es queda como texto."""
    if v is None or isinstance(v, (int, float)):
        return v
    try:
        return float(v)
    except (TypeError, ValueError):
        return str(v)


def columnas_de_bd(
    filas: list[dict[str, Any]],
    analitos: list[dict],
    ingredientes: set[str] | None = None,
) -> list[tuple[str, list[Columna]]]:
    """Grupos de columnas de la descarga, ya acotados a lo que tiene datos.

    `ingredientes`: si viene, solo esos analitos (por el código con el que están
    en la base); `None` = todos los que aparezcan.
    """
    permitidos = {codigo_de_formato(i) for i in ingredientes} if ingredientes else None
    con_resultado: dict[str, str] = {}   # código de formato -> nombre (por si cae en «Otros»)
    con_dosis: set[str] = set()
    con_solicitado: set[str] = set()
    for fila in filas:
        for cod in fila.get("solicitados") or ():
            if permitidos is None or cod in permitidos:
                con_solicitado.add(cod)
        for codigo, res in (fila.get("resultados") or {}).items():
            cod = codigo_de_formato(codigo)
            if permitidos is not None and cod not in permitidos:
                continue
            con_resultado.setdefault(cod, res.get("nombre") or codigo)
        for codigo in fila.get("dosis") or {}:
            cod = codigo_de_formato(codigo)
            if permitidos is None or cod in permitidos:
                con_dosis.add(cod)

    etiquetas_fila = {etiqueta: clave for clave, etiqueta in CAMPOS_FUNGICIDAS}
    grupos: list[tuple[str, list[Columna]]] = [
        ("GENERAL", [("general", clave, etiqueta) for clave, etiqueta in GENERALES_BD])
    ]
    conocidos: set[str] = set()
    for titulo, columnas in _grupos_exportacion(analitos)[1:]:  # [0] es el GENERAL de Solicitudes
        elegidas: list[Columna] = []
        pares: list[Columna] = []       # fungicidas: «X Solicitado» y «X Dosis» de cada analito
        resultados: list[Columna] = []  # fungicidas: los resultados, juntos tras los pares
        es_fungicidas = titulo.startswith(_PREFIJO_FUNGICIDAS)
        for tipo, clave, etiqueta in columnas:
            if tipo == "analito":
                conocidos.add(clave)
                if es_fungicidas:
                    # Un analito pedido sigue teniendo su columna de resultado, aunque aún no llegue.
                    if clave in con_resultado or clave in con_solicitado or clave in con_dosis:
                        resultados.append((tipo, clave, etiqueta))
                elif clave in con_resultado:
                    elegidas.append((tipo, clave, etiqueta))
            elif tipo == "analito_dosis":
                if es_fungicidas:
                    if clave in con_dosis or clave in con_solicitado or clave in con_resultado:
                        pares.append(("analito_solicitado", clave, f"{clave} Solicitado"))
                        pares.append((tipo, clave, etiqueta))
                elif clave in con_dosis:
                    elegidas.append((tipo, clave, etiqueta))
            elif tipo == "campo" and etiqueta not in _CAMPOS_OMITIDOS:
                # Tipo Aplicación, Gasto y el ensayo van siempre que el grupo tenga
                # analitos: son las mismas columnas que la base «con muestra».
                if not any(t in ("analito", "analito_dosis") for t, _, _ in [*elegidas, *pares, *resultados]):
                    continue
                if etiqueta == "Tipo Aplicación":
                    elegidas.append(("tipo_aplicacion", clave, etiqueta))
                elif etiqueta in etiquetas_fila:
                    elegidas.append(("dato_fila", etiquetas_fila[etiqueta], etiqueta))
        elegidas = [*elegidas, *pares, *resultados]
        if elegidas and es_fungicidas:
            codigos = {clave for _, clave, _ in elegidas}
            labs: set[str] = set()
            for fila in filas:
                tiene = any(
                    codigo_de_formato(c) in codigos
                    for c in [*(fila.get("resultados") or {}), *(fila.get("dosis") or {}), *(fila.get("solicitados") or ())]
                )
                if tiene:
                    labs |= _laboratorios_de(fila.get("laboratorio"))
            titulo = titulo_fungicidas(labs)
        if elegidas:
            grupos.append((titulo, elegidas))

    otros: list[Columna] = []
    for cod in sorted(set(con_resultado) - conocidos):
        otros.append(("analito", cod, con_resultado[cod] if con_resultado[cod] == cod else f"{con_resultado[cod]} ({cod})"))
    for cod in sorted(con_dosis - conocidos):
        otros.append(("analito_dosis", cod, f"{cod} Dosis"))
    if otros:
        grupos.append((TITULO_OTROS, otros))
    return grupos


def construir_workbook_bd(
    filas: list[dict[str, Any]],
    analitos: list[dict],
    ingredientes: set[str] | None = None,
    nota_filtro: str | None = None,
) -> Workbook:
    """`filas`: una por solicitud, con las claves de GENERALES_BD y además
    `resultados` {código: {"valor": ..., "nombre": ...}}, `dosis` {código: valor}
    y `tipo_aplicacion`."""
    filas = expandir_por_posicion(
        filas,
        tipo=lambda f: f.get("tipo_aplicacion"),
        posicion=lambda f: f.get("posicion_muestreo"),
        con_posicion=lambda f, p: {**f, "posicion_muestreo": p},
    )
    completar_solicitados(filas, analitos)
    wb = Workbook()
    ws = wb.active
    ws.title = "BD"
    grupos = columnas_de_bd(filas, analitos, ingredientes)
    columnas = [c for _, grupo in grupos for c in grupo]

    inicio = 1
    for titulo, grupo in grupos:
        fin = inicio + len(grupo) - 1
        if fin > inicio:
            ws.merge_cells(start_row=1, start_column=inicio, end_row=1, end_column=fin)
        celda = ws.cell(row=1, column=inicio, value=titulo)
        celda.font = Font(bold=True, size=10, color=VERDE_OSCURO)
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for col in range(inicio, fin + 1):
            ws.cell(row=1, column=col).fill = PatternFill("solid", fgColor=VERDE_CLARO)
            ws.cell(row=1, column=col).border = _BORDE_COMPLETO
        inicio = fin + 1

    for col, (_, _, etiqueta) in enumerate(columnas, start=1):
        c = ws.cell(row=2, column=col, value=etiqueta)
        c.font = Font(bold=True, size=9, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=VERDE_OSCURO)
        c.border = _BORDE_COMPLETO
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 28
    ws.row_dimensions[2].height = 36

    for fila_idx, fila in enumerate(filas, start=3):
        resultados = {codigo_de_formato(k): v for k, v in (fila.get("resultados") or {}).items()}
        dosis = {codigo_de_formato(k): v for k, v in (fila.get("dosis") or {}).items()}
        for col, (tipo, clave, _) in enumerate(columnas, start=1):
            if tipo == "general":
                valor = _valor_celda(fila.get(clave))
            elif tipo == "analito":
                valor = _numero((resultados.get(clave) or {}).get("valor"))
            elif tipo == "analito_dosis":
                valor = _numero(dosis.get(clave))
            elif tipo == "analito_solicitado":
                valor = "✓" if clave in (fila.get("solicitados") or ()) else None
            elif tipo == "dato_fila":  # Gasto, Código de Ensayo, N° Ensayo
                valor = _valor_celda(fila.get(clave))
            else:  # tipo_aplicacion
                valor = fila.get("tipo_aplicacion")
            celda = ws.cell(row=fila_idx, column=col, value=valor if valor not in (None, "") else None)
            celda.border = _BORDE_COMPLETO
            es_dato = tipo in ("analito", "analito_dosis", "analito_solicitado")
            celda.alignment = Alignment(horizontal="center" if es_dato else "left", vertical="center")
            if tipo == "analito" and valor is not None:
                celda.font = Font(bold=True, color=VERDE_OSCURO)
            if isinstance(valor, date):
                celda.number_format = "DD-MM-YYYY"
        ws.row_dimensions[fila_idx].height = 22

    ws.freeze_panes = "A3"
    ultima = max(2, 2 + len(filas))
    ws.auto_filter.ref = f"A2:{ws.cell(row=ultima, column=len(columnas)).coordinate}"
    ws.sheet_view.showGridLines = False
    for col in range(1, len(columnas) + 1):
        letra = ws.cell(row=2, column=col).column_letter
        etiqueta = str(ws.cell(row=2, column=col).value or "")
        ws.column_dimensions[letra].width = min(28, max(13, len(etiqueta) * 0.85))

    if nota_filtro:
        # Hoja aparte para que quien reciba el archivo sepa que NO es la BD completa.
        nota = wb.create_sheet("Filtros aplicados")
        nota["A1"] = "Esta descarga NO es la base de datos completa"
        nota["A1"].font = Font(bold=True, color="B45309", size=12)
        nota["A2"] = nota_filtro
        nota["A2"].alignment = Alignment(wrap_text=True, vertical="top")
        nota.column_dimensions["A"].width = 90
    return wb
