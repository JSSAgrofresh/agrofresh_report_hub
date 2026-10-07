"""
Listas de distribución de resultados: exportar, comparar con un Excel y aplicar
SOLO lo que el admin general confirma.

Es lo mismo que hacían los scripts `importar_contactos_resultado.py` y
`auditar_contactos_resultado.py`, pero desde Administración General: sin tocar
el servidor y sin subir "el archivo completo" para un cambio chico.

Flujo:
  1. Exportar: el Excel trae, planta por planta, lo que el sistema tiene HOY
     (Admin Report Hub, Comercial, Técnico y los correos del cliente por especie).
  2. Se edita el Excel (agregar o corregir correos) y se sube. También sirve la
     hoja «Informes Laboratorios-Pack Line» del Excel maestro.
  3. El sistema compara y devuelve una lista de CAMBIOS (+ agregar, − quitar).
     No escribe nada.
  4. Se confirman los que se quieren y recién ahí se aplican, con respaldo.

Reglas:
  - Comercial → copia (cc). Técnico y Admin Report Hub → copia oculta (bcc).
  - Una celda vacía NO quita a nadie. Para sacar a alguien se deja la celda con
    los demás correos y sin el suyo. (Así una columna sin datos en el Excel no
    borra listas ya cargadas.)
  - Las plantas se reconocen por su Sold To + Ship To ignorando mayúsculas,
    tildes y espacios repetidos; al escribir se respeta el nombre del sistema.
  - Cada tipo de servicio tiene SU lista (`servicio`, ver servicios.py): la de
    siempre es Línea de proceso; Actimist es otra, con su propio listado de
    plantas. Todo endpoint recibe `servicio` (vacío = Línea de proceso) y solo
    lee y escribe los contactos de ese servicio: guardar en uno nunca toca el
    otro, aunque una planta se llame igual en los dos.
"""
from __future__ import annotations

import difflib
import io
import logging
import re
import unicodedata
from datetime import datetime
from typing import Any

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from fastapi.responses import Response
from pydantic import BaseModel, Field

from . import config_store
from .auth import Usuario, solo_admin_general
from .listados import clave_normalizada as _clave_esp
from .servicios import clave_servicio, es_del_servicio, es_servicio_con_listado, tablas

logger = logging.getLogger(__name__)

router = APIRouter(prefix="/api/listas-distribucion", tags=["listas-distribucion"])

ARCHIVO_CONTACTOS = "contactos_laboratorio.json"
LAB_COMPARTIDO = "AGROFRESH"
HOJA_MAESTRO = "Informes Laboratorios-Pack Line"

# Columna del Excel → especies del sistema.
CATEGORIA_ESPECIES: dict[str, list[str]] = {
    "Manzana y Pera": ["Manzana", "Pera"],
    "Kiwi": ["Kiwi"],
    "Carozos": ["Durazno", "Nectarina"],
    "Cerezas": ["Cereza"],
    "Citricos": ["Clementina", "Limón", "Mandarina", "Naranja", "Pomelo"],
    "Arandanos": ["Arándano"],
    "Paltas": ["Palta"],
    "Nueces y Pasas": ["Nueces", "Pasas"],
    "Granada": ["Granada"],
}
CATEGORIAS = tuple(CATEGORIA_ESPECIES)

# rol → (cargo que se guarda, a qué copia va)
ROLES: dict[str, tuple[str, str]] = {
    "admin": ("Admin", "bcc"),
    "comercial": ("Comercial", "cc"),
    "tecnico": ("Técnico", "bcc"),
}
CAMPOS_INTERNOS = tuple(ROLES)
ETIQUETA_CAMPO = {
    "admin": "Admin Report Hub · copia oculta",
    "comercial": "Comercial a cargo · copia",
    "tecnico": "Técnico a cargo · copia oculta",
    **{c: f"Correos del cliente · {c}" for c in CATEGORIAS},
}

_PLACEHOLDERS = {"", "na", "n/a", "#n/a", "none", "-", "--", "falta", "ok", "no aplica", "nan"}
_EMAIL_RE = re.compile(r"^[^@\s;,]+@[^@\s;,]+\.[^@\s;,]+$")


# ---------------------------------------------------------------------------
# Normalización
# ---------------------------------------------------------------------------

def norm(texto: str | None) -> str:
    """Sin tildes, sin mayúsculas y con los espacios colapsados."""
    t = unicodedata.normalize("NFKD", str(texto or ""))
    t = "".join(c for c in t if not unicodedata.combining(c))
    return re.sub(r"\s+", " ", t).strip().casefold()


def clave_planta(sold_to: str | None, ship_to: str | None) -> tuple[str, str]:
    return norm(sold_to), norm(ship_to)


def limpiar_emails(valor: Any) -> tuple[list[str], list[str]]:
    """(correos válidos sin repetir, textos que no son correo).

    Separa por «;», «,» o salto de línea y quita «mailto:». Los marcadores de
    «sin dato» (NA, #N/A, FALTA…) se ignoran sin avisar.
    """
    if valor is None:
        return [], []
    validos: list[str] = []
    invalidos: list[str] = []
    vistos: set[str] = set()
    for parte in re.split(r"[;,\n]", str(valor)):
        p = re.sub(r"^mailto:", "", parte.strip(), flags=re.IGNORECASE).strip()
        if p.casefold() in _PLACEHOLDERS:
            continue
        if _EMAIL_RE.match(p):
            p = p.lower()
            if p not in vistos:
                vistos.add(p)
                validos.append(p)
        else:
            invalidos.append(p)
    return validos, invalidos


def rol_de(contacto: dict) -> str:
    """admin / comercial / tecnico de un contacto interno ya cargado."""
    cargo = norm(contacto.get("cargo"))
    if "admin" in cargo:
        return "admin"
    if "tecnic" in cargo:
        return "tecnico"
    if "comercial" in cargo:
        return "comercial"
    return "tecnico" if contacto.get("tipo_copia") == "bcc" else "comercial"


_CATEGORIA_DE_ESPECIE = {
    _clave_esp(esp): cat for cat, lista in CATEGORIA_ESPECIES.items() for esp in lista
}


def categoria_de_especie(especie: str | None) -> str | None:
    return _CATEGORIA_DE_ESPECIE.get(_clave_esp(especie or ""))


# ---------------------------------------------------------------------------
# Lo que el sistema tiene hoy
# ---------------------------------------------------------------------------

def _fila_vacia(sold_to: str, ship_to: str) -> dict:
    return {
        "sold_to": sold_to, "ship_to": ship_to,
        "admin": [], "comercial": [], "tecnico": [],
        "clientes": {c: [] for c in CATEGORIAS},
        "copia_mal": [],
    }


def _agregar(lista: list[str], email: str) -> None:
    if email.casefold() not in {e.casefold() for e in lista}:
        lista.append(email)


def del_servicio(contactos: list[dict], servicio: str = "") -> list[dict]:
    """Solo los contactos de la lista de ese servicio (vacío = Línea de proceso)."""
    return [c for c in contactos if es_del_servicio(c, servicio)]


def estado_desde_contactos(contactos: list[dict]) -> dict[tuple[str, str], dict]:
    """Una fila por planta con lo que HOY recibe cada rol (solo contactos activos)."""
    plantas: dict[tuple[str, str], dict] = {}
    for c in sorted(contactos, key=lambda c: c.get("orden", 0)):
        if c.get("tipo") not in ("resultado_cliente", "resultado_interno"):
            continue
        email = str(c.get("email") or "").strip()
        if not email or not c.get("activo", True):
            continue
        sold_to = (c.get("sold_to") or "").strip()
        ship_to = (c.get("ship_to") or "").strip()
        fila = plantas.setdefault(clave_planta(sold_to, ship_to), _fila_vacia(sold_to, ship_to))
        if c["tipo"] == "resultado_interno":
            rol = rol_de(c)
            _agregar(fila[rol], email)
            if (c.get("tipo_copia") or "cc") != ROLES[rol][1]:
                _agregar(fila["copia_mal"], email)
        else:
            especie = (c.get("especie") or "").strip()
            if not especie:
                for cat in CATEGORIAS:
                    _agregar(fila["clientes"][cat], email)
            elif (cat := categoria_de_especie(especie)) is not None:
                _agregar(fila["clientes"][cat], email)
    return plantas


# ---------------------------------------------------------------------------
# Excel → filas
# ---------------------------------------------------------------------------

def leer_filas_excel(contenido: bytes) -> tuple[list[dict], list[str]]:
    """(filas, avisos) del Excel subido. Reconoce las columnas por su encabezado."""
    import openpyxl

    try:
        wb = openpyxl.load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError("No se pudo abrir el archivo: tiene que ser un Excel (.xlsx).") from exc
    hoja = next((ws for ws in wb.worksheets if ws.title == HOJA_MAESTRO), wb.worksheets[0])

    filas_excel = hoja.iter_rows(values_only=True)
    columnas: dict[str, int] = {}
    for fila in filas_excel:
        encabezados = [norm(c) for c in fila]
        if "sold to name" in encabezados and "ship to name" in encabezados:
            for i, h in enumerate(encabezados):
                if h == "sold to name":
                    columnas["sold_to"] = i
                elif h == "ship to name":
                    columnas["ship_to"] = i
                elif "admin" in h:
                    columnas["admin"] = i
                elif "comercial" in h:
                    columnas["comercial"] = i
                elif "tecnic" in h:
                    columnas["tecnico"] = i
                elif "vigente" in h:
                    columnas["vigente"] = i
                elif h in ("s number", "sold to number"):
                    columnas["codigo_sold"] = i
                elif h == "ship to number":
                    columnas["codigo_ship"] = i
                else:
                    for cat in CATEGORIAS:
                        if h == norm(cat):
                            columnas[cat] = i
            break
    if "sold_to" not in columnas:
        raise ValueError(
            "No encontré las columnas «SOLD TO NAME» y «SHIP TO NAME». Usa el Excel que "
            "exporta esta misma pantalla o la hoja «Informes Laboratorios-Pack Line»."
        )

    avisos: list[str] = []
    vistos: set[tuple[str, str]] = set()
    filas: list[dict] = []
    for fila in filas_excel:
        def celda(clave: str) -> Any:
            i = columnas.get(clave)
            return fila[i] if i is not None and i < len(fila) else None

        sold_to = str(celda("sold_to") or "").strip()
        ship_to = str(celda("ship_to") or "").strip()
        if not sold_to and not ship_to:
            continue
        if "vigente" in columnas and norm(str(celda("vigente") or "")) == "no":
            continue
        clave = clave_planta(sold_to, ship_to)
        if clave in vistos:
            avisos.append(f"«{ship_to}» (Sold To «{sold_to}») aparece más de una vez: se usó la primera fila.")
            continue
        vistos.add(clave)

        def correos(campo: str) -> list[str]:
            validos, invalidos = limpiar_emails(celda(campo))
            for texto in invalidos:
                avisos.append(f"«{ship_to}»: «{texto}» no es un correo ({campo}); se ignoró.")
            return validos

        filas.append({
            "sold_to": sold_to, "ship_to": ship_to,
            "codigo_sold": str(celda("codigo_sold") or "").strip() or None,
            "codigo_ship": str(celda("codigo_ship") or "").strip() or None,
            "admin": correos("admin"), "comercial": correos("comercial"), "tecnico": correos("tecnico"),
            "clientes": {cat: correos(cat) for cat in CATEGORIAS},
        })
    return filas, avisos


# ---------------------------------------------------------------------------
# Comparar
# ---------------------------------------------------------------------------

def _sin_los_de(lista: list[str], otros: list[str]) -> list[str]:
    ya = {e.casefold() for e in otros}
    return [e for e in lista if e.casefold() not in ya]


def parecidas(sold_to: str, ship_to: str, listados: dict[tuple[str, str], tuple[str, str]], n: int = 3) -> list[dict]:
    """Hasta `n` plantas de Listados con un nombre parecido, para decirle a quien
    revisa «¿será esta?» cuando el Excel trae un nombre que Listados no tiene."""
    ship_n, sold_n = norm(ship_to), norm(sold_to)
    puntuadas: list[tuple[float, tuple[str, str]]] = []
    for (sold_l, ship_l), oficial in listados.items():
        r_ship = difflib.SequenceMatcher(None, ship_n, ship_l).ratio()
        if r_ship < 0.6:
            continue
        r_sold = difflib.SequenceMatcher(None, sold_n, sold_l).ratio()
        puntuadas.append((r_ship * 0.7 + r_sold * 0.3, oficial))
    puntuadas.sort(key=lambda p: -p[0])
    return [{"sold_to": a, "ship_to": b} for _, (a, b) in puntuadas[:n]]


def comparar(
    estado: dict[tuple[str, str], dict],
    filas: list[dict],
    listados: dict[tuple[str, str], tuple[str, str]] | None = None,
) -> dict:
    """Los cambios que habría que hacer para que el sistema quede como el Excel.

    Una celda vacía no quita a nadie. Cada cambio es independiente: se puede
    confirmar uno y no otro.
    """
    cambios: list[dict] = []
    sin_cambios = 0
    for fila in filas:
        clave = clave_planta(fila["sold_to"], fila["ship_to"])
        actual = estado.get(clave)
        if actual is None:
            if not any(fila[c] for c in CAMPOS_INTERNOS) and not any(fila["clientes"].values()):
                sin_cambios += 1  # fila en blanco: nada que agregar
                continue
            sold_to, ship_to, aviso, sugerencias = fila["sold_to"], fila["ship_to"], None, []
            if listados is not None:
                if clave in listados:
                    sold_to, ship_to = listados[clave]
                else:
                    aviso = "No existe en Listados con ese nombre: la app no la encontrará hasta que coincida."
                    sugerencias = parecidas(sold_to, ship_to, listados)
            cambios.append({
                "id": f"{'|'.join(clave)}|nueva", "tipo": "planta_nueva",
                "planta": {"sold_to": sold_to, "ship_to": ship_to},
                "campo": "planta", "etiqueta": "Planta nueva en las listas",
                "agregar": [], "quitar": [], "corregir": [], "aviso": aviso, "sugerencias": sugerencias,
                "fila": {**fila, "sold_to": sold_to, "ship_to": ship_to},
            })
            continue

        planta = {"sold_to": actual["sold_to"], "ship_to": actual["ship_to"]}
        n_antes = len(cambios)
        for campo in (*CAMPOS_INTERNOS, *CATEGORIAS):
            nuevo = fila[campo] if campo in CAMPOS_INTERNOS else fila["clientes"][campo]
            viejo = actual[campo] if campo in CAMPOS_INTERNOS else actual["clientes"][campo]
            if not nuevo:
                continue  # celda vacía = sin cambios
            agregar, quitar = _sin_los_de(nuevo, viejo), _sin_los_de(viejo, nuevo)
            if agregar or quitar:
                cambios.append({
                    "id": f"{'|'.join(clave)}|{campo}", "tipo": "campo", "planta": planta,
                    "campo": campo, "etiqueta": ETIQUETA_CAMPO[campo],
                    "agregar": agregar, "quitar": quitar, "corregir": [], "aviso": None, "fila": None,
                })
        if actual["copia_mal"]:
            cambios.append({
                "id": f"{'|'.join(clave)}|copia", "tipo": "copia", "planta": planta,
                "campo": "copia", "etiqueta": "Ajustar a copia / copia oculta",
                "agregar": [], "quitar": [], "corregir": list(actual["copia_mal"]),
                "aviso": "El comercial va en copia; el técnico y el admin, en copia oculta.", "fila": None,
            })
        if len(cambios) == n_antes:
            sin_cambios += 1

    claves_excel = {clave_planta(f["sold_to"], f["ship_to"]) for f in filas}
    solo_sistema = [f"{e['ship_to']} ({e['sold_to']})" for k, e in estado.items() if k not in claves_excel]
    return {
        "cambios": cambios,
        "resumen": {
            "plantas_excel": len(filas),
            "plantas_sin_cambios": sin_cambios,
            "plantas_con_cambios": len({c["id"].rsplit("|", 1)[0] for c in cambios}),
            "cambios": len(cambios),
            "plantas_solo_sistema": len(solo_sistema),
            "solo_sistema": sorted(solo_sistema)[:50],
        },
    }


# ---------------------------------------------------------------------------
# Aplicar
# ---------------------------------------------------------------------------

def _contacto(id_: int, sold_to: str, ship_to: str, especie: str, email: str, tipo: str,
              cargo: str, copia: str, orden: int, nombre: str | None = None, servicio: str = "") -> dict:
    contacto = {
        "id": id_, "laboratorio": LAB_COMPARTIDO, "nombre": nombre or email, "email": email,
        "cargo": cargo, "tipo": tipo, "sold_to": sold_to, "ship_to": ship_to, "especie": especie,
        "tipo_copia": copia, "activo": True, "orden": orden,
    }
    # Los de Línea de proceso quedan como siempre (sin la llave); los de
    # Actimist llevan su marca.
    if es_servicio_con_listado(servicio):
        contacto["servicio"] = clave_servicio(servicio)
    return contacto


def _contactos_clientes(
    sold_to: str, ship_to: str, mapa: dict[str, list[str]], sig_id: int, orden0: int,
    nombres: dict[str, dict], servicio: str = "",
) -> tuple[list[dict], int]:
    """Los contactos de cliente de un mapa categoría → correos.

    Si las nueve categorías tienen exactamente los mismos correos se guarda UN
    grupo general (especie vacía: vale para toda especie). Si no, un grupo por
    especie, solo donde hay correos: una especie sin lista no hereda la de otra.
    """
    salida: list[dict] = []
    llenas = {c: mapa.get(c, []) for c in CATEGORIAS}
    claves = [tuple(sorted(e.casefold() for e in v)) for v in llenas.values()]
    general = all(llenas.values()) and len(set(claves)) == 1
    orden = orden0

    def nuevo(especie: str, email: str) -> None:
        nonlocal sig_id, orden
        previo = nombres.get(email.casefold(), {})
        c = _contacto(sig_id, sold_to, ship_to, especie, email, "resultado_cliente",
                      previo.get("cargo", ""), "cc", orden, previo.get("nombre"), servicio)
        salida.append(c)
        sig_id += 1
        orden += 1

    if general:
        for email in llenas[CATEGORIAS[0]]:
            nuevo("", email)
    else:
        for cat in CATEGORIAS:
            for especie in CATEGORIA_ESPECIES[cat]:
                for email in llenas[cat]:
                    nuevo(especie, email)
    return salida, sig_id


def aplicar(contactos: list[dict], cambios: list[dict], servicio: str = "") -> tuple[list[dict], dict]:
    """Aplica SOLO los cambios confirmados sobre una copia de `contactos`.

    Es tolerante: agregar a quien ya está, o quitar a quien ya no está, se
    ignora. Devuelve la lista nueva y un resumen de lo hecho.

    `contactos` es la configuración COMPLETA (todos los servicios), pero solo
    se tocan los del `servicio` pedido: los de la otra lista quedan intactos.
    """
    servicio = clave_servicio(servicio)
    nuevos = [dict(c) for c in contactos]
    sig_id = max((c.get("id", 0) for c in nuevos), default=0) + 1
    hechos = {"aplicados": 0, "plantas": set(), "ignorados": []}

    por_planta: dict[tuple[str, str], list[dict]] = {}
    for c in cambios:
        p = c.get("planta") or {}
        por_planta.setdefault(clave_planta(p.get("sold_to"), p.get("ship_to")), []).append(c)

    for clave, items in por_planta.items():
        def de_la_planta(c: dict) -> bool:
            return c.get("tipo") in ("resultado_cliente", "resultado_interno") and \
                es_del_servicio(c, servicio) and \
                clave_planta(c.get("sold_to"), c.get("ship_to")) == clave

        existentes = [c for c in nuevos if de_la_planta(c)]
        sold_to = existentes[0].get("sold_to", "") if existentes else (items[0]["planta"].get("sold_to") or "")
        ship_to = existentes[0].get("ship_to", "") if existentes else (items[0]["planta"].get("ship_to") or "")
        orden0 = max((c.get("orden", 0) for c in existentes), default=0) + 1
        nombres = {str(c.get("email") or "").casefold(): c for c in existentes}
        mapa_clientes: dict[str, list[str]] | None = None

        for it in items:
            tipo, campo = it.get("tipo"), it.get("campo")
            agregar = [e for e in (it.get("agregar") or []) if _EMAIL_RE.match(str(e))]
            quitar = {str(e).casefold() for e in (it.get("quitar") or [])}
            corregir = {str(e).casefold() for e in (it.get("corregir") or [])}

            if tipo == "planta_nueva":
                if existentes:
                    hechos["ignorados"].append(f"{ship_to}: ya tiene listas; no se creó de nuevo.")
                    continue
                fila = it.get("fila") or {}
                for rol, (cargo, copia) in ROLES.items():
                    for email in limpiar_emails(";".join(fila.get(rol) or []))[0]:
                        nuevos.append(_contacto(sig_id, sold_to, ship_to, "", email, "resultado_interno", cargo, copia,
                                                orden0, servicio=servicio))
                        sig_id += 1
                        orden0 += 1
                cli, sig_id = _contactos_clientes(
                    sold_to, ship_to, {c: (fila.get("clientes") or {}).get(c, []) for c in CATEGORIAS},
                    sig_id, orden0, {}, servicio)
                nuevos.extend(cli)
                hechos["aplicados"] += 1
                hechos["plantas"].add(clave)

            elif tipo == "campo" and campo in CAMPOS_INTERNOS:
                cargo, copia = ROLES[campo]
                nuevos[:] = [
                    c for c in nuevos
                    if not (de_la_planta(c) and c["tipo"] == "resultado_interno" and rol_de(c) == campo
                            and str(c.get("email") or "").casefold() in quitar)
                ]
                ya = {str(c.get("email") or "").casefold() for c in nuevos
                      if de_la_planta(c) and c["tipo"] == "resultado_interno" and rol_de(c) == campo
                      and c.get("activo", True)}
                for email in agregar:
                    if email.casefold() in ya:
                        continue
                    nuevos.append(_contacto(sig_id, sold_to, ship_to, "", email.lower(), "resultado_interno", cargo, copia,
                                            orden0, servicio=servicio))
                    sig_id += 1
                    orden0 += 1
                hechos["aplicados"] += 1
                hechos["plantas"].add(clave)

            elif tipo == "campo" and campo in CATEGORIAS:
                if mapa_clientes is None:
                    mapa_clientes = {c: list(v) for c, v in
                                     (estado_desde_contactos(existentes).get(clave) or _fila_vacia("", ""))["clientes"].items()}
                mapa_clientes[campo] = [e for e in mapa_clientes[campo] if e.casefold() not in quitar]
                for email in agregar:
                    _agregar(mapa_clientes[campo], email.lower())
                hechos["aplicados"] += 1
                hechos["plantas"].add(clave)

            elif tipo == "copia":
                for c in nuevos:
                    if de_la_planta(c) and c["tipo"] == "resultado_interno" \
                            and str(c.get("email") or "").casefold() in corregir:
                        c["tipo_copia"] = ROLES[rol_de(c)][1]
                hechos["aplicados"] += 1
                hechos["plantas"].add(clave)
            else:
                hechos["ignorados"].append(f"{ship_to}: cambio desconocido ({tipo}/{campo}).")

        if mapa_clientes is not None:
            # Se reemplazan los contactos de cliente activos que caen en una categoría;
            # los desactivados y los de especies fuera de las categorías no se tocan.
            def reemplazable(c: dict) -> bool:
                if not (de_la_planta(c) and c["tipo"] == "resultado_cliente" and c.get("activo", True)):
                    return False
                esp = (c.get("especie") or "").strip()
                return not esp or categoria_de_especie(esp) is not None

            nuevos[:] = [c for c in nuevos if not reemplazable(c)]
            cli, sig_id = _contactos_clientes(sold_to, ship_to, mapa_clientes, sig_id, orden0, nombres, servicio)
            nuevos.extend(cli)

    hechos["plantas"] = len(hechos["plantas"])
    return nuevos, hechos


# ---------------------------------------------------------------------------
# Exportar
# ---------------------------------------------------------------------------

ENCABEZADOS_EXCEL = ["#", "SOLD TO NAME", "SHIP TO NAME", "Admin Report Hub", "Comercial a cargo",
                     "Técnico a cargo", *CATEGORIAS]


def construir_excel(estado: dict[tuple[str, str], dict], vacias: list[tuple[str, str]] | None = None) -> bytes:
    """El Excel con las listas de hoy; `vacias` agrega plantas de Listados sin lista."""
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = HOJA_MAESTRO
    ws.append(ENCABEZADOS_EXCEL)
    filas = sorted(estado.values(), key=lambda f: (norm(f["sold_to"]), norm(f["ship_to"])))
    n = 0
    for f in filas:
        n += 1
        ws.append([n, f["sold_to"], f["ship_to"], "; ".join(f["admin"]), "; ".join(f["comercial"]),
                   "; ".join(f["tecnico"]), *("; ".join(f["clientes"][c]) for c in CATEGORIAS)])
    for sold_to, ship_to in sorted(vacias or [], key=lambda p: (norm(p[0]), norm(p[1]))):
        n += 1
        ws.append([n, sold_to, ship_to])
    for celda in ws[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="3F7D20")
        celda.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "D2"
    for col, ancho in zip("ABCDEFGHIJKLMNOP", (6, 34, 40, 34, 28, 28, *([30] * 9))):
        ws.column_dimensions[col].width = ancho
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Estado para la tabla y alta en Listados
# ---------------------------------------------------------------------------

def estado_para_tabla(
    estado: dict[tuple[str, str], dict],
    listados: dict[tuple[str, str], tuple[str, str]] | None,
    incluir_sin_lista: bool = False,
) -> dict:
    """Las filas de la tabla dinámica: una por planta, con lo que recibe cada rol."""
    filas: list[dict] = []
    for clave, f in estado.items():
        filas.append({
            "sold_to": f["sold_to"], "ship_to": f["ship_to"],
            "admin": f["admin"], "comercial": f["comercial"], "tecnico": f["tecnico"],
            "clientes": f["clientes"],
            # qué correos de cada rol están en la copia equivocada (para marcarlos en la celda)
            "copia_mal": {rol: [e for e in f[rol] if e in f["copia_mal"]] for rol in CAMPOS_INTERNOS},
            "en_listados": None if listados is None else clave in listados,
            "sin_contactos": False,
        })
    sin_lista = 0
    if listados is not None:
        faltan = [par for k, par in listados.items() if k not in estado]
        sin_lista = len(faltan)
        if incluir_sin_lista:
            for sold_to, ship_to in faltan:
                f = _fila_vacia(sold_to, ship_to)
                filas.append({
                    "sold_to": sold_to, "ship_to": ship_to, "admin": [], "comercial": [], "tecnico": [],
                    "clientes": f["clientes"], "copia_mal": {rol: [] for rol in CAMPOS_INTERNOS},
                    "en_listados": True, "sin_contactos": True,
                })
    filas.sort(key=lambda f: (norm(f["sold_to"]), norm(f["ship_to"])))
    return {
        "filas": filas,
        "resumen": {
            "plantas_con_lista": len(estado),
            "plantas_listados": None if listados is None else len(listados),
            "listados_sin_lista": None if listados is None else sin_lista,
        },
    }


def asegurar_planta(cur, sold_to: str, ship_to: str, codigo_sold: str | None = None,
                    codigo_ship: str | None = None, servicio: str = "") -> dict:
    """Crea en Listados el cliente (Sold To) y la planta (Ship To) si no existen.

    Compara sin mayúsculas, tildes ni espacios repetidos para no duplicar. Devuelve
    los nombres oficiales (los que ya había, o los escritos) y qué se creó.
    Escribe en el listado del `servicio` (Actimist tiene el suyo).
    """
    t_cliente, t_planta = tablas(servicio)
    sold_to, ship_to = re.sub(r"\s+", " ", sold_to).strip(), re.sub(r"\s+", " ", ship_to).strip()
    if not sold_to or not ship_to:
        raise ValueError("Falta el Sold To o el Ship To.")
    cur.execute(f"SELECT id, nombre FROM {t_cliente}")
    cliente = next((c for c in cur.fetchall() if norm(c["nombre"]) == norm(sold_to)), None)
    cliente_creado = cliente is None
    if cliente is None:
        cur.execute(f"INSERT INTO {t_cliente} (nombre, codigo_sap, activo) VALUES (%s, %s, true) RETURNING id",
                    (sold_to, codigo_sold))
        cliente = {"id": cur.fetchone()["id"], "nombre": sold_to}
    cur.execute(f"SELECT id, nombre FROM {t_planta} WHERE cliente_id = %s", (cliente["id"],))
    planta = next((p for p in cur.fetchall() if norm(p["nombre"]) == norm(ship_to)), None)
    planta_creada = planta is None
    if planta is None:
        cur.execute(f"INSERT INTO {t_planta} (cliente_id, nombre, codigo_sap, activo) VALUES (%s, %s, %s, true) RETURNING id",
                    (cliente["id"], ship_to, codigo_ship))
        planta = {"id": cur.fetchone()["id"], "nombre": ship_to}
    return {"sold_to": cliente["nombre"], "ship_to": planta["nombre"],
            "cliente_creado": cliente_creado, "planta_creada": planta_creada}


# ---------------------------------------------------------------------------
# Endpoints (solo admin general)
# ---------------------------------------------------------------------------

def _listados(servicio: str = "") -> dict[tuple[str, str], tuple[str, str]] | None:
    """Plantas de Listados: clave normalizada → (Sold To, Ship To) oficiales.
    Cada servicio tiene su listado (Actimist: migración 0049)."""
    t_cliente, t_planta = tablas(servicio)
    try:
        from .db import conexion, cursor_dict

        with conexion(escribir=False) as conn, cursor_dict(conn) as cur:
            cur.execute(
                f"SELECT c.nombre AS cliente, p.nombre AS planta FROM {t_planta} p "
                f"JOIN {t_cliente} c ON c.id = p.cliente_id WHERE p.activo AND c.activo"
            )
            return {clave_planta(f["cliente"], f["planta"]): (f["cliente"], f["planta"]) for f in cur.fetchall()}
    except Exception:
        logger.warning("No se pudo leer Listados para las listas de distribución.", exc_info=True)
        return None


def _clientes_listados(servicio: str = "") -> list[str]:
    t_cliente, _ = tablas(servicio)
    try:
        from .db import conexion, cursor_dict

        with conexion(escribir=False) as conn, cursor_dict(conn) as cur:
            cur.execute(f"SELECT nombre FROM {t_cliente} WHERE activo ORDER BY nombre")
            return [f["nombre"] for f in cur.fetchall()]
    except Exception:
        logger.warning("No se pudo leer los clientes de Listados.", exc_info=True)
        return []


class CambiosIn(BaseModel):
    cambios: list[dict] = Field(default_factory=list, max_length=2000)


@router.get("/estado")
def estado_actual(
    sin_lista: bool = False, servicio: str = "", _: Usuario = Depends(solo_admin_general)
) -> dict:
    """Lo que el sistema tiene hoy, una fila por planta: la base de la tabla dinámica."""
    servicio = clave_servicio(servicio)
    estado = estado_desde_contactos(del_servicio(config_store.leer(ARCHIVO_CONTACTOS, []), servicio))
    resultado = estado_para_tabla(estado, _listados(servicio), sin_lista)
    resultado["clientes"] = _clientes_listados(servicio)
    resultado["servicio"] = servicio
    return resultado


@router.get("/excel")
def exportar(todas: bool = False, servicio: str = "", _: Usuario = Depends(solo_admin_general)) -> Response:
    servicio = clave_servicio(servicio)
    estado = estado_desde_contactos(del_servicio(config_store.leer(ARCHIVO_CONTACTOS, []), servicio))
    vacias: list[tuple[str, str]] = []
    if todas:
        lis = _listados(servicio) or {}
        vacias = [par for k, par in lis.items() if k not in estado]
    sufijo = f"_{servicio}" if es_servicio_con_listado(servicio) else ""
    nombre = f"listas_distribucion{sufijo}_{datetime.now():%Y-%m-%d}.xlsx"
    return Response(
        construir_excel(estado, vacias),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=\"{nombre}\"; filename*=UTF-8''{nombre}"},
    )


@router.post("/comparar")
async def comparar_excel(
    archivo: UploadFile = File(...), servicio: str = "", _: Usuario = Depends(solo_admin_general)
) -> dict:
    servicio = clave_servicio(servicio)
    contenido = await archivo.read()
    if len(contenido) > 15 * 1024 * 1024:
        raise HTTPException(413, "El archivo pesa más de 15 MB.")
    try:
        filas, avisos = leer_filas_excel(contenido)
    except ValueError as exc:
        raise HTTPException(400, str(exc)) from exc
    resultado = comparar(
        estado_desde_contactos(del_servicio(config_store.leer(ARCHIVO_CONTACTOS, []), servicio)),
        filas, _listados(servicio),
    )
    resultado["resumen"]["avisos"] = avisos[:100]
    return resultado


@router.post("/aplicar")
def aplicar_cambios(datos: CambiosIn, servicio: str = "", usuario: Usuario = Depends(solo_admin_general)) -> dict:
    servicio = clave_servicio(servicio)
    if not datos.cambios:
        raise HTTPException(400, "No hay cambios confirmados para aplicar.")
    creados = {"clientes": 0, "plantas": 0}
    for it in datos.cambios:
        if it.get("tipo") == "planta_nueva" and it.get("crear_en_listados"):
            fila = it.get("fila") or {}
            try:
                from .db import conexion, cursor_dict

                with conexion() as conn, cursor_dict(conn) as cur:
                    r = asegurar_planta(cur, (it.get("planta") or {}).get("sold_to", ""),
                                        (it.get("planta") or {}).get("ship_to", ""),
                                        fila.get("codigo_sold"), fila.get("codigo_ship"), servicio)
            except ValueError as exc:
                raise HTTPException(400, str(exc)) from exc
            except Exception as exc:
                logger.exception("No se pudo crear la planta en Listados")
                raise HTTPException(500, f"No se pudo crear la planta en Listados: {exc}") from exc
            it["planta"] = {"sold_to": r["sold_to"], "ship_to": r["ship_to"]}
            it["fila"] = {**fila, "sold_to": r["sold_to"], "ship_to": r["ship_to"]}
            creados["clientes"] += r["cliente_creado"]
            creados["plantas"] += r["planta_creada"]
    actuales = config_store.leer(ARCHIVO_CONTACTOS, [])
    respaldo = f"contactos_laboratorio_respaldo_{datetime.now():%Y%m%d_%H%M%S}.json"
    config_store.escribir(respaldo, actuales)
    nuevos, hechos = aplicar(actuales, datos.cambios, servicio)
    config_store.escribir(ARCHIVO_CONTACTOS, nuevos)
    logger.info("Listas de distribución (%s): %s aplicó %d cambios en %d plantas (respaldo %s)",
                servicio or "linea", usuario.email, hechos["aplicados"], hechos["plantas"], respaldo)
    return {**hechos, "respaldo": respaldo, "listados_creados": creados}
