"""Plan de entregables de los comerciales («Plan Servicio de Análisis & Entrega de Información»).

Cada comercial lleva un Excel: una fila por cliente (planta) y servicio, y una columna por semana
con cuántos análisis se prometen (W41, W42… como semanas ISO; W01 y W02 son del año siguiente).
Este módulo solo LEE esos Excel y los deja en una tabla larga (cliente × servicio × semana);
no escribe nada. Es la base de la futura carga a la base de datos.

Los totales de los Excel no son confiables (columnas que se saltan una semana, totales escritos
a mano, fórmulas que suman filas de más). Aquí manda la suma de las semanas, y cuando una fila
dice otra cosa se avisa en `descuadres`.
"""
from __future__ import annotations

import re
import unicodedata
from datetime import date, timedelta
from typing import Any

import openpyxl

SERVICIOS = ("Residuales", "Patologías", "Microbiología", "Entrega de registros", "FSMA")


def sin_tildes(texto: str) -> str:
    return "".join(c for c in unicodedata.normalize("NFD", str(texto)) if unicodedata.category(c) != "Mn")


def categoria_servicio(servicio: str) -> str:
    """«Patologias L1», «Fitopatológico», «Microbilogicos»… → una de `SERVICIOS`."""
    s = sin_tildes(servicio).lower().strip()
    if s.startswith("residual"):
        return "Residuales"
    if s.startswith(("patolog", "fitopatolog")):
        return "Patologías"
    if s.startswith("microbi"):
        return "Microbiología"
    if s.startswith("registro"):
        return "Entrega de registros"
    if s.startswith("fsma"):
        return "FSMA"
    return str(servicio).strip()


def lunes_de_semana(codigo: str) -> date:
    """«W41» → lunes de esa semana ISO en la temporada (W01…W19 son del año siguiente)."""
    n = int(codigo[1:])
    return date.fromisocalendar(2027 if n < 20 else 2026, n, 1)


def comercial_del_titulo(titulo: str) -> str:
    """«Plan Servicio … 2026-27  -  JC Labarca:» → «JC Labarca»."""
    return re.split(r"\s+-\s+", str(titulo).strip())[-1].strip(" :")


def _fila_de_cabecera(ws) -> int:
    for r in range(1, 15):
        if str(ws.cell(r, 2).value or "").strip().upper() == "CLIENTE":
            return r
    raise ValueError("No encuentro la fila de cabecera («CLIENTE»).")


def leer_plan(ruta_o_buffer: Any, nombre_archivo: str = "") -> dict[str, Any]:
    """Lee un Excel de plan. Devuelve `comercial`, `filas` (una por cliente × servicio × semana),
    `semanas` y `descuadres` (filas cuyo TOTAL no coincide con la suma de sus semanas)."""
    ws = openpyxl.load_workbook(ruta_o_buffer, data_only=True).worksheets[0]
    cab = _fila_de_cabecera(ws)
    comercial = comercial_del_titulo(ws.cell(2, 2).value or "")
    semanas: dict[int, str] = {}
    col_total = None
    for c in ws[cab]:
        v = str(c.value or "").strip()
        if re.fullmatch(r"W\d\d", v):
            semanas[c.column] = v
        elif v.upper() == "TOTAL":
            col_total = c.column
            break
    if not semanas:
        raise ValueError("No encuentro las columnas de semanas (W41, W42…).")

    filas: list[dict[str, Any]] = []
    descuadres: list[dict[str, Any]] = []
    for r in range(cab + 2, ws.max_row + 1):
        cliente, servicio = ws.cell(r, 2).value, ws.cell(r, 3).value
        if cliente is None and servicio is None:
            continue
        if str(cliente or "").strip().upper() == "TOTAL":
            break  # lo que sigue son los totales y la simbología
        if cliente is None or not isinstance(servicio, str):
            continue
        cliente_txt, servicio_txt = str(cliente).strip(), servicio.strip()
        linea = ""
        m = re.search(r"\bL(\d)\b", servicio_txt)
        if m:
            linea = f"L{m.group(1)}"
        planta = cliente_txt
        mc = re.fullmatch(r"(.*?)\s+L(\d)", cliente_txt)  # en el plan de Suarez la línea va pegada al cliente
        if mc:
            planta, linea = mc.group(1).strip(), f"L{mc.group(2)}"
        suma = 0
        for col, sem in semanas.items():
            v = ws.cell(r, col).value
            if v in (None, ""):
                continue
            v = int(v)
            suma += v
            desde = lunes_de_semana(sem)
            filas.append({
                "comercial": comercial, "archivo": nombre_archivo, "cliente_plan": planta, "linea": linea,
                "servicio_original": servicio_txt, "servicio": categoria_servicio(servicio_txt), "semana": sem,
                "semana_desde": desde.isoformat(), "semana_hasta": (desde + timedelta(days=6)).isoformat(),
                "cantidad": v, "fila_excel": r,
            })
        total = ws.cell(r, col_total).value if col_total else None
        if total is not None and int(total) != suma:
            descuadres.append({"fila_excel": r, "cliente_plan": planta, "servicio": servicio_txt, "total_excel": int(total), "suma_semanas": suma})
    return {"comercial": comercial, "filas": filas, "semanas": list(semanas.values()), "descuadres": descuadres}


# ── el Excel unificado (una tabla larga con todos los comerciales) ──

HOJA_UNIFICADA = "Entregables"
_COLUMNAS_UNIFICADAS = ("Comercial", "Cliente (como en el plan)", "Línea", "Servicio", "Servicio (original)", "Semana", "Desde", "Hasta", "Cantidad", "Archivo", "Fila en el Excel")


def _es_unificado(wb) -> bool:
    if HOJA_UNIFICADA not in wb.sheetnames:
        return False
    fila = [str(c.value or "").strip() for c in wb[HOJA_UNIFICADA][1]]
    return tuple(fila[: len(_COLUMNAS_UNIFICADAS)]) == _COLUMNAS_UNIFICADAS


def leer_planes(ruta_o_buffer: Any, nombre_archivo: str = "") -> list[dict[str, Any]]:
    """Lee un Excel de plan y devuelve UN plan por comercial: el de cada comercial por separado
    (un plan) o el Excel unificado de todos (uno por comercial que trae)."""
    wb = openpyxl.load_workbook(ruta_o_buffer, data_only=True)
    if not _es_unificado(wb):
        if hasattr(ruta_o_buffer, "seek"):
            ruta_o_buffer.seek(0)
        return [leer_plan(ruta_o_buffer, nombre_archivo)]
    por_comercial: dict[str, dict[str, Any]] = {}
    for f in wb[HOJA_UNIFICADA].iter_rows(min_row=2, values_only=True):
        if not f[0] or not f[1] or f[8] in (None, ""):
            continue
        comercial = str(f[0]).strip()
        plan = por_comercial.setdefault(comercial, {"comercial": comercial, "filas": [], "semanas": [], "descuadres": []})
        semana = str(f[5]).strip()
        if semana not in plan["semanas"]:
            plan["semanas"].append(semana)
        plan["filas"].append({
            "comercial": comercial, "archivo": str(f[9] or nombre_archivo), "cliente_plan": str(f[1]).strip(), "linea": str(f[2] or ""),
            "servicio_original": str(f[4] or ""), "servicio": str(f[3] or ""), "semana": semana,
            "semana_desde": str(f[6] or ""), "semana_hasta": str(f[7] or ""), "cantidad": int(f[8]), "fila_excel": f[10],
        })
    return list(por_comercial.values())
