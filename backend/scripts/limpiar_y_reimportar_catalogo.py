"""
Limpia el catálogo de clientes (sold_to) y plantas (ship_to) y lo reimporta
desde el Excel maestro de contactos.

Uso (desde la carpeta backend):
  .venv\\Scripts\\python.exe scripts\\limpiar_y_reimportar_catalogo.py ruta\\al\\archivo.xlsx
  .venv\\Scripts\\python.exe scripts\\limpiar_y_reimportar_catalogo.py ruta\\al\\archivo.xlsx --aplicar

Qué hace:
  1. Pone a NULL planta_id en todas las solicitudes (FK nullable).
  2. Borra todos los registros de mapeo_confirmado (referencian cliente).
  3. Trunca las tablas planta y cliente.
  4. Reimporta clientes y plantas desde el Excel (solo filas VIGENTE POST VENTA = SI).

Las solicitudes NO se borran. Sus datos (sold_to_raw, ship_to_raw, formulario)
quedan intactos — solo pierden la FK numérica a planta, que ya no se usa.
"""

import sys
from pathlib import Path

SCRIPT_DIR = Path(__file__).resolve().parent
BACKEND_DIR = SCRIPT_DIR.parent
sys.path.insert(0, str(BACKEND_DIR))

DEFAULT_EXCEL = SCRIPT_DIR / "importar_contactos_resultado_excel.xlsx"


def es_vigente(valor) -> bool:
    return str(valor or "").strip().upper() == "SI"


def leer_pares(path: str) -> list[tuple[str, str]]:
    try:
        import openpyxl
    except ImportError:
        sys.exit("Falta openpyxl: pip install openpyxl")

    wb = openpyxl.load_workbook(path, data_only=True)
    hoja = "Informes Laboratorios-Pack Line"
    if hoja not in wb.sheetnames:
        sys.exit(f"No se encontró la hoja '{hoja}' en {path}")
    ws = wb[hoja]

    pares: list[tuple[str, str]] = []
    vistos: set[tuple[str, str]] = set()

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=True):
        if not row[0]:
            continue
        if not es_vigente(row[2]):
            continue
        sold_to = str(row[4] or "").strip()
        ship_to = str(row[6] or "").strip()
        if not sold_to or not ship_to:
            continue
        key = (sold_to, ship_to)
        if key not in vistos:
            vistos.add(key)
            pares.append(key)

    return pares


def main():
    aplicar = "--aplicar" in sys.argv
    excel_arg = next((a for a in sys.argv[1:] if not a.startswith("-")), None)
    excel_path = excel_arg or str(DEFAULT_EXCEL)

    import os
    if not os.path.exists(excel_path):
        sys.exit(
            f"No se encontró el Excel en:\n  {excel_path}\n"
            f"Pásalo como argumento o ponlo en:\n  {DEFAULT_EXCEL}"
        )

    print(f"Leyendo: {excel_path}")
    pares = leer_pares(excel_path)

    sold_tos = sorted({p[0] for p in pares})
    ship_tos_por_cliente: dict[str, list[str]] = {}
    for sold, ship in pares:
        ship_tos_por_cliente.setdefault(sold, []).append(ship)

    print()
    print("=" * 60)
    print("RESUMEN EXCEL")
    print("=" * 60)
    print(f"Sold To únicos:   {len(sold_tos)}")
    print(f"Pares únicos:     {len(pares)}")
    print()

    from app.db import conexion

    with conexion(escribir=aplicar) as conn:
        cur = conn.cursor()

        # Conteos actuales
        cur.execute("SELECT COUNT(*) FROM cliente")
        n_clientes = cur.fetchone()[0]
        cur.execute("SELECT COUNT(*) FROM planta")
        n_plantas = cur.fetchone()[0]
        cur.execute("SELECT COUNT(*) FROM mapeo_confirmado")
        n_mapeos = cur.fetchone()[0]
        cur.execute("SELECT COUNT(*) FROM solicitud WHERE planta_id IS NOT NULL")
        n_sol_con_planta = cur.fetchone()[0]

        print(f"Estado actual en BD:")
        print(f"  Clientes (sold_to):   {n_clientes}")
        print(f"  Plantas (ship_to):    {n_plantas}")
        print(f"  Mapeos confirmados:   {n_mapeos}  → se borran")
        print(f"  Solicitudes con FK:   {n_sol_con_planta}  → quedan en NULL (datos intactos)")
        print()
        print(f"Después de reimportar:")
        print(f"  Clientes (sold_to):   {len(sold_tos)}")
        print(f"  Plantas (ship_to):    {len(pares)}")
        print()

        if not aplicar:
            print(">> DRY RUN — no se escribió nada.")
            print(">> Agrega --aplicar para ejecutar.")
            return

        # 1. Soltar FK de solicitudes
        cur.execute("UPDATE solicitud SET planta_id = NULL WHERE planta_id IS NOT NULL")
        print(f"  ✓ {cur.rowcount} solicitudes con planta_id → NULL")

        # 2. Borrar mapeos confirmados (referencian cliente)
        cur.execute("DELETE FROM mapeo_confirmado")
        print(f"  ✓ {cur.rowcount} mapeos confirmados eliminados")

        # 3. Truncar planta y cliente (en orden por FK)
        cur.execute("DELETE FROM planta")
        print(f"  ✓ Plantas eliminadas")
        cur.execute("DELETE FROM cliente")
        print(f"  ✓ Clientes eliminados")

        # Reiniciar secuencias para IDs limpios
        cur.execute("ALTER SEQUENCE cliente_id_seq RESTART WITH 1")
        cur.execute("ALTER SEQUENCE planta_id_seq RESTART WITH 1")

        # 4. Insertar clientes
        for nombre in sold_tos:
            cur.execute(
                "INSERT INTO cliente (nombre) VALUES (%s) ON CONFLICT (nombre) DO NOTHING",
                (nombre,),
            )
        print(f"  ✓ {len(sold_tos)} clientes insertados")

        # Mapa nombre → id
        cur.execute("SELECT id, nombre FROM cliente")
        cliente_id_por_nombre = {r[1]: r[0] for r in cur.fetchall()}

        # 5. Insertar plantas
        insertadas = 0
        for sold, ships in ship_tos_por_cliente.items():
            cid = cliente_id_por_nombre.get(sold)
            if cid is None:
                print(f"  AVISO: cliente '{sold}' no encontrado, saltando sus plantas")
                continue
            for ship in ships:
                cur.execute(
                    "INSERT INTO planta (cliente_id, nombre) VALUES (%s, %s) ON CONFLICT (cliente_id, nombre) DO NOTHING",
                    (cid, ship),
                )
                insertadas += 1
        print(f"  ✓ {insertadas} plantas insertadas")

        print()
        print("✓ Catálogo actualizado exitosamente.")


if __name__ == "__main__":
    main()
