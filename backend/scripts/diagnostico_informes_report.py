"""
Diagnóstico de TODOS los informes que están en Report: ¿están bien colocados y bien parseados?

SOLO LEE: no cambia nada en la base ni en R2. Compara cada informe con su solicitud de Toma de muestras (servicio,
laboratorio, Sold To, Ship To, especie, fecha, analitos pedidos), revisa su coherencia interna (datos que faltan, fechas,
resultados vacíos o raros) y, con --pdf, vuelve a leer su PDF con los lectores del Converter (Quiteca y AgroFresh) y
compara OT, fecha, especie y cada resultado con lo guardado.

Deja un Excel en logs/ (Resumen, Por laboratorio, Hallazgos con todo el detalle) y un resumen en pantalla.

Uso (desde la carpeta backend):
    .venv\\Scripts\\python.exe scripts\\diagnostico_informes_report.py                 # rápido, sin leer PDF
    .venv\\Scripts\\python.exe scripts\\diagnostico_informes_report.py --pdf           # además relee los PDF (más lento)
    .venv\\Scripts\\python.exe scripts\\diagnostico_informes_report.py --lab Quiteca   # un laboratorio
    .venv\\Scripts\\python.exe scripts\\diagnostico_informes_report.py --desde 2026-10-01
"""
from __future__ import annotations

import argparse
import sys
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path

_BACKEND = Path(__file__).resolve().parent.parent
if str(_BACKEND) not in sys.path:
    sys.path.insert(0, str(_BACKEND))

from app import diagnostico_informes as di  # noqa: E402
from app.db import conexion, cursor_dict  # noqa: E402

_SQL_INFORMES = """
    SELECT s.id, s.nro_solicitud, s.laboratorio, to_jsonb(s) ->> 'referencia' AS referencia, to_jsonb(s) ->> 'servicio' AS servicio,
           s.tipo_servicio, s.especie, s.variedad, s.planta_id, s.fecha_muestreo, s.fecha_entrada,
           s.fecha_recepcion, s.fecha_informe, s.fecha_analisis,
           COALESCE(c.nombre, s.sold_to_raw) AS sold_to, COALESCE(p.nombre, s.ship_to_raw) AS ship_to
    FROM solicitud s
    LEFT JOIN planta p ON p.id = s.planta_id
    LEFT JOIN cliente c ON c.id = p.cliente_id
    WHERE s.vigente {filtros}
    ORDER BY s.id
"""
_SQL_RESULTADOS = """
    SELECT r.solicitud_id, r.analito_id, r.analito_raw, a.codigo, r.valor_num, r.valor_texto
    FROM resultado r LEFT JOIN analito a ON a.id = r.analito_id
"""
_SQL_TOMA = """
    SELECT upper(btrim(numero_solicitud)) AS ot, laboratorio, sold_to, ship_to, especie, fecha_muestreo, datos
    FROM solicitud_archivo WHERE btrim(COALESCE(numero_solicitud, '')) <> ''
"""
_SQL_PDF = "SELECT btrim(nro_informe) AS nro, lower(btrim(laboratorio)) AS lab, r2_key FROM informe_auditoria ORDER BY subido_en"


def _leer_pdf(cache: dict, nro: str, laboratorio: str):
    """{datos, resultados} del PDF de ese informe leído con el lector del Converter, o None si no se puede."""
    from app import correo_informes_homogeneo as hom, r2_auditoria as r2a

    clave = cache.get((str(nro or "").strip(), di.lab_base(laboratorio)))
    lab = "AgroFresh" if di.lab_base(laboratorio) == "agrofresh" else "Quiteca" if di.lab_base(laboratorio) == "quiteca" else None
    if not clave or not lab or not r2a.disponible():
        return None
    try:
        datos = r2a.descargar(clave)
        return hom.leer(datos, lab) if datos else None
    except Exception:  # noqa: BLE001 — un PDF raro no frena el diagnóstico
        return None


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--lab", default="", help="Solo este laboratorio (Quiteca, Agrofresh, ALS, Diagnofruit)")
    ap.add_argument("--desde", default="", help="Solo informes con muestreo desde esta fecha (AAAA-MM-DD)")
    ap.add_argument("--pdf", action="store_true", help="Vuelve a leer el PDF de cada informe de Quiteca y AgroFresh y lo compara")
    ap.add_argument("--salida", default="", help="Archivo Excel de salida (por defecto logs/diagnostico_informes_<fecha>.xlsx)")
    args = ap.parse_args(argv)

    filtros, params = "", {}
    if args.lab:
        filtros += " AND lower(split_part(s.laboratorio, '(', 1)) LIKE %(lab)s"
        params["lab"] = f"%{args.lab.strip().lower()}%"
    if args.desde:
        filtros += " AND s.fecha_muestreo >= %(desde)s"
        params["desde"] = args.desde

    with conexion(escribir=False) as conn, cursor_dict(conn) as cur:
        cur.execute(_SQL_INFORMES.format(filtros=filtros), params)
        informes = cur.fetchall()
        cur.execute(_SQL_RESULTADOS)
        resultados: dict[int, list[dict]] = defaultdict(list)
        for r in cur.fetchall():
            resultados[r["solicitud_id"]].append(r)
        cur.execute(_SQL_TOMA)
        tomas = {t["ot"]: t for t in cur.fetchall()}
        pdfs: dict = {}
        if args.pdf:
            try:
                cur.execute(_SQL_PDF)
                pdfs = {(f["nro"], f["lab"]): f["r2_key"] for f in cur.fetchall()}
            except Exception as exc:  # noqa: BLE001
                print(f"No se pudo leer informe_auditoria ({exc}): sigue sin comparar PDF.")

    filas: list[dict] = []
    por_lab: dict[str, Counter] = defaultdict(Counter)
    total_lab: Counter = Counter()
    limpios = 0
    for n, inf in enumerate(informes, 1):
        ot = str(inf.get("referencia") or "").strip().upper()
        toma = tomas.get(ot) if ot else None
        leido = _leer_pdf(pdfs, inf["nro_solicitud"], inf["laboratorio"]) if args.pdf and di.lab_base(inf["laboratorio"]) in ("agrofresh", "quiteca") else False
        if args.pdf and leido is None and di.lab_base(inf["laboratorio"]) not in ("agrofresh", "quiteca"):
            leido = False
        hallazgos = di.revisar(inf, resultados.get(inf["id"], []), toma, leido)
        lab = inf["laboratorio"] or "(sin laboratorio)"
        total_lab[lab] += 1
        graves = [h for h in hallazgos if h["gravedad"] != di.INFO]
        if not graves:
            limpios += 1
        for h in hallazgos:
            por_lab[lab][h["codigo"]] += 1
            filas.append({"id": inf["id"], "informe": inf["nro_solicitud"], "laboratorio": lab, "ot": inf.get("referencia"), "sold_to": inf.get("sold_to"),
                          "ship_to": inf.get("ship_to"), "especie": inf.get("especie"), "fecha_muestreo": inf.get("fecha_muestreo"),
                          "servicio": di.NOMBRE_SERVICIO.get(di._servicio(inf.get("servicio")), inf.get("servicio")), **h})
        if args.pdf and n % 200 == 0:
            print(f"  … {n} de {len(informes)} informes revisados")

    por_codigo = Counter((f["gravedad"], f["codigo"]) for f in filas)
    informes_por_codigo: dict[str, set] = defaultdict(set)
    for f in filas:
        informes_por_codigo[f["codigo"]].add(f["id"])

    print(f"\nInformes revisados: {len(informes)}  ·  sin hallazgos graves o medios: {limpios}  ·  con algo que mirar: {len(informes) - limpios}")
    print(f"{'Gravedad':8} {'Hallazgo':26} {'Informes':>8}")
    orden = {di.ALTO: 0, di.MEDIO: 1, di.INFO: 2}
    for (gravedad, codigo), _ in sorted(por_codigo.items(), key=lambda kv: (orden[kv[0][0]], -len(informes_por_codigo[kv[0][1]]))):
        print(f"{gravedad:8} {codigo:26} {len(informes_por_codigo[codigo]):>8}")

    salida = Path(args.salida) if args.salida else _BACKEND / "logs" / f"diagnostico_informes_{datetime.now():%Y%m%d_%H%M}.xlsx"
    salida.parent.mkdir(parents=True, exist_ok=True)
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Resumen"
    ws.append(["Gravedad", "Hallazgo", "Informes afectados"])
    for (gravedad, codigo), _ in sorted(por_codigo.items(), key=lambda kv: (orden[kv[0][0]], -len(informes_por_codigo[kv[0][1]]))):
        ws.append([gravedad, codigo, len(informes_por_codigo[codigo])])
    ws.append([])
    ws.append(["Informes revisados", len(informes)])
    ws.append(["Sin hallazgos graves ni medios", limpios])
    ws2 = wb.create_sheet("Por laboratorio")
    codigos = sorted({c for cont in por_lab.values() for c in cont})
    ws2.append(["Laboratorio", "Informes", *codigos])
    for lab in sorted(total_lab):
        ws2.append([lab, total_lab[lab], *[por_lab[lab].get(c, 0) for c in codigos]])
    ws3 = wb.create_sheet("Hallazgos")
    cab = ["Gravedad", "Hallazgo", "N° informe", "Laboratorio", "OT", "Sold To", "Ship To", "Especie", "Fecha muestreo", "Servicio en Report", "Detalle", "Qué hacer", "ID"]
    ws3.append(cab)
    for f in sorted(filas, key=lambda f: (orden[f["gravedad"]], f["codigo"], str(f["informe"]))):
        ws3.append([f["gravedad"], f["codigo"], f["informe"], f["laboratorio"], f["ot"], f["sold_to"], f["ship_to"], f["especie"], f["fecha_muestreo"],
                    f["servicio"], f["detalle"], f["sugerencia"], f["id"]])
    for hoja in (ws, ws2, ws3):
        for c in hoja[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="3F6B2A")
        for i in range(1, hoja.max_column + 1):
            hoja.column_dimensions[get_column_letter(i)].width = 26
    ws3.auto_filter.ref = ws3.dimensions
    ws3.freeze_panes = "A2"
    ws3.column_dimensions["K"].width = 90
    wb.save(salida)
    print(f"\nExcel con todo el detalle: {salida}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
