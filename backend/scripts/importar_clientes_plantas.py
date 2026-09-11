"""
Importa el catálogo de Sold To (clientes) y Ship To (plantas) desde el Excel
maestro de contactos de resultado.

Uso (desde la carpeta backend):
  .venv\\Scripts\\python.exe scripts\\importar_clientes_plantas.py ruta\\al\\archivo.xlsx
  .venv\\Scripts\\python.exe scripts\\importar_clientes_plantas.py ruta\\al\\archivo.xlsx --aplicar

Lee la hoja "Informes Laboratorios-Pack Line" del Excel:
  - Col E → nombre del cliente (sold_to)
  - Col G → nombre de la planta (ship_to)

Comportamiento:
  - Omite filas sin sold_to o ship_to.
  - Omite filas con VIGENTE POST VENTA ≠ "SI" (col C).
  - Hace upsert seguro: inserta clientes/plantas que faltan, no toca los existentes.
  - --aplicar para escribir; sin él solo muestra qué haría (dry run).
"""

import sys
from pathlib import Path

SCRIPT_DIR = Path(__file__).resolve().parent
BACKEND_DIR = SCRIPT_DIR.parent
sys.path.insert(0, str(BACKEND_DIR))

DEFAULT_EXCEL = SCRIPT_DIR / "importar_contactos_resultado_excel.xlsx"


def es_vigente(valor) -> bool:
    return str(valor or "").strip().upper() == "SI"


def leer_pares(path: str) -> list[tuple[str, str]]:
    try:
        import openpyxl
    except ImportError:
        sys.exit("Falta openpyxl: pip install openpyxl")

    wb = openpyxl.load_workbook(path, data_only=True)
    hoja = "Informes Laboratorios-Pack Line"
    if hoja not in wb.sheetnames:
        sys.exit(f"No se encontró la hoja '{hoja}' en {path}")
    ws = wb[hoja]

    pares: list[tuple[str, str]] = []
    vistos: set[tuple[str, str]] = set()

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=True):
        if not row[0]:
            continue
        if not es_vigente(row[2]):
            continue
        sold_to = str(row[4] or "").strip()
        ship_to = str(row[6] or "").strip()
        if not sold_to or not ship_to:
            continue
        key = (sold_to, ship_to)
        if key not in vistos:
            vistos.add(key)
            pares.append(key)

    return pares


def main():
    aplicar = "--aplicar" in sys.argv
    excel_arg = next((a for a in sys.argv[1:] if not a.startswith("-")), None)
    excel_path = excel_arg or str(DEFAULT_EXCEL)

    import os
    if not os.path.exists(excel_path):
        sys.exit(
            f"No se encontró el Excel en:\n  {excel_path}\n"
            "Pásalo como argumento o ponlo en:\n"
            f"  {DEFAULT_EXCEL}"
        )

    print(f"Leyendo: {excel_path}")
    pares = leer_pares(excel_path)

    sold_tos = sorted({p[0] for p in pares})
    ship_tos_por_cliente: dict[str, list[str]] = {}
    for sold, ship in pares:
        ship_tos_por_cliente.setdefault(sold, []).append(ship)

    print()
    print("=" * 60)
    print("RESUMEN EXCEL")
    print("=" * 60)
    print(f"Sold To únicos:   {len(sold_tos)}")
    print(f"Pares únicos:     {len(pares)}")
    print()

    # Conectar a la BD
    from app.db import conexion

    with conexion(escribir=aplicar) as conn:
        cur = conn.cursor()

        # Clientes existentes
        cur.execute("SELECT id, nombre FROM cliente")
        existentes_cliente = {r[1]: r[0] for r in cur.fetchall()}

        # Plantas existentes (cliente_id, nombre) → id
        cur.execute("SELECT id, cliente_id, nombre FROM planta")
        existentes_planta = {(r[1], r[2]): r[0] for r in cur.fetchall()}

        nuevos_clientes: list[str] = []
        nuevas_plantas: list[tuple[str, str]] = []

        for sold in sold_tos:
            if sold not in existentes_cliente:
                nuevos_clientes.append(sold)

        for sold, ships in ship_tos_por_cliente.items():
            # El cliente puede ser nuevo (aún sin id); lo marcamos igual
            cliente_id = existentes_cliente.get(sold, -1)
            for ship in ships:
                if (cliente_id, ship) not in existentes_planta:
                    nuevas_plantas.append((sold, ship))

        print(f"Clientes a insertar: {len(nuevos_clientes)}")
        print(f"Plantas a insertar:  {len(nuevas_plantas)}")
        print()

        if nuevos_clientes:
            print("Clientes nuevos:")
            for c in nuevos_clientes[:20]:
                print(f"  + {c}")
            if len(nuevos_clientes) > 20:
                print(f"  ... y {len(nuevos_clientes) - 20} más.")
            print()

        if nuevas_plantas:
            print("Plantas nuevas (sold_to · ship_to):")
            for s, p in nuevas_plantas[:20]:
                print(f"  + {s} · {p}")
            if len(nuevas_plantas) > 20:
                print(f"  ... y {len(nuevas_plantas) - 20} más.")
            print()

        if not aplicar:
            print(">> DRY RUN — no se escribió nada.")
            print(">> Agrega --aplicar para guardar.")
            return

        # Insertar clientes nuevos
        for nombre in nuevos_clientes:
            cur.execute(
                "INSERT INTO cliente (nombre) VALUES (%s) ON CONFLICT (nombre) DO NOTHING",
                (nombre,),
            )

        # Refrescar mapa de clientes (incluye los recién insertados)
        cur.execute("SELECT id, nombre FROM cliente")
        cliente_id_por_nombre = {r[1]: r[0] for r in cur.fetchall()}

        # Insertar plantas nuevas
        for sold, ship in nuevas_plantas:
            cid = cliente_id_por_nombre.get(sold)
            if cid is None:
                print(f"  AVISO: no se encontró cliente '{sold}', saltando planta '{ship}'")
                continue
            cur.execute(
                """
                INSERT INTO planta (cliente_id, nombre)
                VALUES (%s, %s)
                ON CONFLICT (cliente_id, nombre) DO NOTHING
                """,
                (cid, ship),
            )

        print(f"✓ Insertados {len(nuevos_clientes)} clientes y {len(nuevas_plantas)} plantas.")


if __name__ == "__main__":
    main()
