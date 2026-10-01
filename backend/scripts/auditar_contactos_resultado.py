"""
Audita los contactos de resultados: Excel maestro vs lo que hay HOY en el sistema.

Solo lee: no escribe nada en la base, en R2 ni en el Excel.

Compara, planta por planta (hoja «Informes Laboratorios-Pack Line»):
  - Para (clientes), CC (comerciales), CCO (técnicos y admin Report Hub);
  - que los técnicos estén en CCO y los comerciales en CC;
  - que el nombre de la planta exista tal cual en Listados (la búsqueda es por texto exacto);
  - lo que de verdad sale hoy en un correo (la misma función que usa la app).

Genera un Excel en logs/ con el detalle y un resumen por pantalla.

Uso (desde la carpeta backend):
  .venv\\Scripts\\python.exe scripts\\auditar_contactos_resultado.py "C:\\ruta\\Master cliente - plantas SAP.xlsx"
"""
from __future__ import annotations

import os
import re
import sys
from datetime import datetime
from pathlib import Path

BACKEND_DIR = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(BACKEND_DIR))

from scripts import importar_contactos_resultado as imp  # noqa: E402

CARPETA_LOGS = os.path.join(BACKEND_DIR.parent, "logs")
ESPECIE_MUESTRA = "Manzana"  # con la que se simula «lo que sale hoy»


def _norm(txt: str) -> str:
    return re.sub(r"\s+", " ", str(txt or "")).strip().casefold()


def leer_filas_excel(ws) -> list[dict]:
    """Una fila por planta vigente, con sus correos separados por rol."""
    filas: list[dict] = []
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=True):
        if not row[0] or not imp.es_vigente(row[2]):
            continue
        clientes: list[str] = []
        for col in imp.COLS_ESP:
            for e in imp.limpiar_emails(row[col]):
                if e not in clientes:
                    clientes.append(e)
        comercial = imp.limpiar_emails(row[8])
        filas.append({
            "sold_to": str(row[4] or "").strip(),
            "ship_to": str(row[6] or "").strip(),
            "clientes": clientes,
            "comercial": comercial,
            "tecnico": imp.limpiar_emails(row[9]),
            "admin": imp.limpiar_emails(row[7]),
            "pendiente_andres": imp.COMERCIAL_EXCLUIDO in comercial,
        })
    return filas


def auditar_planta(fila: dict, contactos: list[dict]) -> list[str]:
    """Problemas de UNA planta: lo que dice el Excel contra lo que hay en `contactos`."""
    st, sh = fila["sold_to"].strip(), fila["ship_to"].strip()
    pool = [
        c for c in contactos
        if c.get("tipo") in ("resultado_cliente", "resultado_interno")
        and (c.get("sold_to") or "").strip() == st and (c.get("ship_to") or "").strip() == sh
    ]
    problemas: list[str] = []
    if fila["pendiente_andres"]:
        problemas.append("PENDIENTE: comercial es Andrés González (el importador salta la fila)")
    if not pool:
        problemas.append("SIN CONTACTOS EN EL SISTEMA para esta planta")
        return problemas

    activos = [c for c in pool if c.get("activo", True) and c.get("email")]
    cc = {c["email"].strip().casefold() for c in activos
          if c["tipo"] == "resultado_interno" and c.get("tipo_copia") != "bcc"}
    bcc = {c["email"].strip().casefold() for c in activos
           if c["tipo"] == "resultado_interno" and c.get("tipo_copia") == "bcc"}
    para = {c["email"].strip().casefold() for c in activos if c["tipo"] == "resultado_cliente"}

    for e in fila["tecnico"]:
        k = e.casefold()
        if k not in cc and k not in bcc:
            problemas.append(f"TÉCNICO FALTA: {e}")
        elif k not in bcc:
            problemas.append(f"TÉCNICO EN CC (debe ir en CCO): {e}")
    for e in fila["comercial"]:
        k = e.casefold()
        if k not in cc and k not in bcc:
            problemas.append(f"COMERCIAL FALTA: {e}")
        elif k not in cc:
            problemas.append(f"COMERCIAL EN CCO (debe ir en CC): {e}")
    for e in fila["admin"]:
        if e.casefold() not in bcc:
            problemas.append(f"ADMIN FALTA en CCO: {e}")

    esperados_internos = {e.casefold() for e in (*fila["comercial"], *fila["tecnico"], *fila["admin"])}
    for k in sorted((cc | bcc) - esperados_internos):
        problemas.append(f"SOBRA EN SISTEMA (no está en el Excel): {k}")

    excel_cli = {e.casefold() for e in fila["clientes"]}
    if excel_cli - para:
        problemas.append("CLIENTES FALTAN en sistema: " + ", ".join(sorted(excel_cli - para)))
    if para - excel_cli:
        problemas.append("CLIENTES SOBRAN en sistema: " + ", ".join(sorted(para - excel_cli)))
    return problemas


def _listados() -> list[tuple[str, str]] | None:
    """(cliente, planta) de Listados, o None si no hay base."""
    try:
        from app.db import conexion, cursor_dict

        with conexion(escribir=False) as conn, cursor_dict(conn) as cur:
            cur.execute("SELECT c.nombre AS cliente, p.nombre AS planta FROM planta p JOIN cliente c ON c.id = p.cliente_id")
            return [(f["cliente"], f["planta"]) for f in cur.fetchall()]
    except Exception as exc:  # sin base no se cae: se omite esta revisión
        print(f"(No se pudo leer Listados: {exc})")
        return None


def estado_nombre(fila: dict, listados: list[tuple[str, str]] | None) -> str:
    if listados is None:
        return "sin revisar"
    if (fila["sold_to"], fila["ship_to"]) in listados:
        return "OK"
    if any(_norm(a) == _norm(fila["sold_to"]) and _norm(b) == _norm(fila["ship_to"]) for a, b in listados):
        return "DISTINTO en espacios/mayúsculas (la app no lo encuentra)"
    return "NO EXISTE en Listados"


def sale_hoy(fila: dict, contactos: list[dict]) -> dict[str, list[str]]:
    """Lo que la app arma de verdad para esta planta, usando su misma función."""
    from app import toma_muestras as tm

    original = tm._leer_config
    tm._leer_config = lambda nombre, defecto: contactos if nombre == "contactos_laboratorio.json" else original(nombre, defecto)
    try:
        return tm.destinatarios_resultado_por_tipo("AGROFRESH", fila["ship_to"], fila["sold_to"], ESPECIE_MUESTRA)
    finally:
        tm._leer_config = original


def main() -> None:
    import openpyxl
    from app import config_store

    ruta = next((a for a in sys.argv[1:] if not a.startswith("-")), None)
    if not ruta or not os.path.exists(ruta):
        sys.exit("Pasa la ruta del Excel maestro como argumento.")

    filas = leer_filas_excel(imp.leer_excel(ruta))
    contactos = config_store.leer("contactos_laboratorio.json", [])
    listados = _listados()
    print(f"Excel: {len(filas)} plantas vigentes · Sistema: {len(contactos)} contactos de laboratorio")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Auditoría por planta"
    ws.append(["Sold To", "Ship To", "Nombre en Listados", "Problemas",
               f"Sale hoy ({ESPECIE_MUESTRA}) · Para", "CC", "CCO", "Excel · Comercial (CC)", "Excel · Técnico (CCO)"])
    conteo: dict[str, int] = {}
    con_problemas = 0
    for f in filas:
        pr = auditar_planta(f, contactos)
        nombre = estado_nombre(f, listados)
        if nombre not in ("OK", "sin revisar"):
            pr.append(f"NOMBRE: {nombre}")
        hoy = sale_hoy(f, contactos)
        for p in pr:
            k = re.split(r"[:(]", p)[0].strip()
            conteo[k] = conteo.get(k, 0) + 1
        con_problemas += bool(pr)
        ws.append([f["sold_to"], f["ship_to"], nombre, "\n".join(pr) or "OK",
                   "; ".join(hoy["to"]), "; ".join(hoy["cc"]), "; ".join(hoy["bcc"]),
                   "; ".join(f["comercial"]), "; ".join(f["tecnico"])])

    excel_claves = {(f["sold_to"], f["ship_to"]) for f in filas}
    ws2 = wb.create_sheet("En sistema, no en Excel")
    ws2.append(["Sold To", "Ship To", "Contactos"])
    visto: dict[tuple[str, str], int] = {}
    for c in contactos:
        if c.get("tipo") in ("resultado_cliente", "resultado_interno"):
            k = ((c.get("sold_to") or "").strip(), (c.get("ship_to") or "").strip())
            if k not in excel_claves:
                visto[k] = visto.get(k, 0) + 1
    for (a, b), n in sorted(visto.items()):
        ws2.append([a, b, n])

    for hoja in (ws, ws2):
        hoja.freeze_panes = "A2"
        for col, ancho in zip("ABCDEFGHI", (34, 40, 26, 60, 40, 40, 40, 32, 32)):
            hoja.column_dimensions[col].width = ancho

    os.makedirs(CARPETA_LOGS, exist_ok=True)
    salida = os.path.join(CARPETA_LOGS, f"auditoria_contactos_{datetime.now():%Y%m%d_%H%M%S}.xlsx")
    wb.save(salida)

    print(f"\nPlantas con algún problema: {con_problemas} de {len(filas)}")
    for k, n in sorted(conteo.items(), key=lambda p: -p[1]):
        print(f"  {n:>4}  {k}")
    print(f"  {len(visto):>4}  plantas en el sistema que no están en el Excel")
    print(f"\nDetalle: {salida}\n")


if __name__ == "__main__":
    main()
