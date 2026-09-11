"""
Importa los contactos de "Resultado a clientes" desde el Excel maestro.

Uso (desde la carpeta backend):
  .venv\\Scripts\\python.exe scripts\\importar_contactos_resultado.py           # preview
  .venv\\Scripts\\python.exe scripts\\importar_contactos_resultado.py --aplicar # escribe

Lee la hoja "Informes Laboratorios-Pack Line" del archivo Excel y genera
entradas en contactos_laboratorio.json organizadas por (sold_to, ship_to,
especie).

Reglas:
- Columna H (Admin Report Hub)  → resultado_interno bcc (copia oculta)
- Columna I (Comercial a cargo) → resultado_interno cc
- Columna J (Técnico a cargo)   → resultado_interno cc (vacío actualmente)
- Columnas K-S (especies)       → resultado_cliente (destinatario directo)
- Filas con comercial = andres.gonzalez@agrofresh.com → se saltan (pendiente)
- Filas con VIGENTE POST VENTA ≠ "SI" → se saltan
- Cuando los correos de cliente son iguales en todas las especies de una fila
  se crea UN solo grupo con especie="" (aplica a todas via fallback).
- Cuando difieren se expanden las categorías del Excel a especies individuales.

Categorías del Excel → especies del sistema:
  Manzana y Pera → Manzana, Pera
  Kiwi           → Kiwi
  Carozos        → Durazno, Nectarina
  Cerezas        → Cereza
  Citricos       → Clementina, Limón, Mandarina, Naranja, Pomelo
  Arandanos      → Arándano
  Paltas         → Palta
  Nueces y Pasas → Nueces, Pasas
  Granada        → Granada
"""

import json
import os
import re
import sys
from pathlib import Path

# ---------------------------------------------------------------------------
# Paths
# ---------------------------------------------------------------------------

SCRIPT_DIR = Path(__file__).resolve().parent
BACKEND_DIR = SCRIPT_DIR.parent
EXCEL_PATH = SCRIPT_DIR / "importar_contactos_resultado_excel.xlsx"

# El config_store usa STORAGE_DIR del .env o la carpeta storage/ por defecto
sys.path.insert(0, str(BACKEND_DIR))
from app import config  # noqa: E402

STORAGE_DIR = config.STORAGE_DIR
JSON_PATH = os.path.join(STORAGE_DIR, "solicitudes", "_config", "contactos_laboratorio.json")

# ---------------------------------------------------------------------------
# Constantes
# ---------------------------------------------------------------------------

COMERCIAL_EXCLUIDO = "andres.gonzalez@agrofresh.com"

# Columna Excel → lista de especies del sistema
CATEGORIA_ESPECIES: dict[str, list[str]] = {
    "Manzana y Pera": ["Manzana", "Pera"],
    "Kiwi":           ["Kiwi"],
    "Carozos":        ["Durazno", "Nectarina"],
    "Cerezas":        ["Cereza"],
    "Citricos":       ["Clementina", "Limón", "Mandarina", "Naranja", "Pomelo"],
    "Arandanos":      ["Arándano"],
    "Paltas":         ["Palta"],
    "Nueces y Pasas": ["Nueces", "Pasas"],
    "Granada":        ["Granada"],
}

# Índice de columna (0-based) → clave normalizada de categoría
COLS_ESP: dict[int, str] = {
    10: "Manzana y Pera",
    11: "Kiwi",
    12: "Carozos",
    13: "Cerezas",
    14: "Citricos",
    15: "Arandanos",
    16: "Paltas",
    17: "Nueces y Pasas",
    18: "Granada",
}

LAB_COMPARTIDO = "AGROFRESH"


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def limpiar_emails(valor) -> list[str]:
    """Extrae correos de una celda: separa por ; o , y quita prefijo mailto:."""
    if not valor:
        return []
    partes = re.split(r"[;,]", str(valor).strip())
    resultado = []
    for p in partes:
        p = re.sub(r"^mailto:", "", p.strip(), flags=re.IGNORECASE).strip()
        if p and "@" in p:
            resultado.append(p.lower())
    return resultado


def es_vigente(valor) -> bool:
    return str(valor or "").strip().upper() == "SI"


# ---------------------------------------------------------------------------
# Lectura del Excel
# ---------------------------------------------------------------------------

def leer_excel(path: str):
    try:
        import openpyxl
    except ImportError:
        sys.exit("Falta openpyxl: pip install openpyxl")

    wb = openpyxl.load_workbook(path, data_only=True)
    hoja = "Informes Laboratorios-Pack Line"
    if hoja not in wb.sheetnames:
        sys.exit(f"No se encontró la hoja '{hoja}' en {path}")
    return wb[hoja]


def construir_grupos(ws) -> list[dict]:
    """Devuelve lista de grupos {sold_to, ship_to, especie, clientes, comercial, admin, tecnico}."""
    grupos: list[dict] = []

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=True):
        if not row[0]:
            continue
        if not es_vigente(row[2]):
            continue

        sold_to  = str(row[4] or "").strip()
        ship_to  = str(row[6] or "").strip()
        admin    = limpiar_emails(row[7])
        comercial = limpiar_emails(row[8])
        tecnico  = limpiar_emails(row[9])

        # Saltar filas de Andres Gonzalez
        if COMERCIAL_EXCLUIDO in comercial:
            continue

        # Correos de clientes por categoría
        cat_emails: dict[str, list[str]] = {}
        for col_idx, cat_nombre in COLS_ESP.items():
            emails = limpiar_emails(row[col_idx])
            if emails:
                cat_emails[cat_nombre] = emails

        if not cat_emails and not comercial and not admin and not tecnico:
            continue

        # ¿Todos los correos de cliente son iguales entre categorías?
        vals = list(cat_emails.values())
        todos_iguales = len(vals) > 0 and all(
            sorted(v) == sorted(vals[0]) for v in vals
        )

        if todos_iguales and vals:
            # Un grupo genérico sin especie (fallback para todas)
            grupos.append({
                "sold_to":  sold_to,
                "ship_to":  ship_to,
                "especie":  "",
                "clientes": vals[0],
                "comercial": comercial,
                "admin":    admin,
                "tecnico":  tecnico,
            })
        else:
            # Un grupo por especie individual
            for cat_nombre, emails in cat_emails.items():
                for especie in CATEGORIA_ESPECIES[cat_nombre]:
                    grupos.append({
                        "sold_to":  sold_to,
                        "ship_to":  ship_to,
                        "especie":  especie,
                        "clientes": emails,
                        "comercial": comercial,
                        "admin":    admin,
                        "tecnico":  tecnico,
                    })

    return grupos


# ---------------------------------------------------------------------------
# Conversión a contactos
# ---------------------------------------------------------------------------

def grupos_a_contactos(grupos: list[dict], id_inicio: int) -> list[dict]:
    contactos: list[dict] = []
    next_id = id_inicio

    for g in grupos:
        sold_to = g["sold_to"]
        ship_to = g["ship_to"]
        especie = g["especie"]
        orden   = 1

        for email in g["clientes"]:
            contactos.append({
                "id":         next_id,
                "laboratorio": LAB_COMPARTIDO,
                "nombre":     email,
                "email":      email,
                "cargo":      "",
                "tipo":       "resultado_cliente",
                "sold_to":    sold_to,
                "ship_to":    ship_to,
                "especie":    especie,
                "tipo_copia": "cc",
                "activo":     True,
                "orden":      orden,
            })
            next_id += 1
            orden   += 1

        for email in g["comercial"]:
            contactos.append({
                "id":         next_id,
                "laboratorio": LAB_COMPARTIDO,
                "nombre":     email,
                "email":      email,
                "cargo":      "Comercial",
                "tipo":       "resultado_interno",
                "sold_to":    sold_to,
                "ship_to":    ship_to,
                "especie":    especie,
                "tipo_copia": "cc",
                "activo":     True,
                "orden":      orden,
            })
            next_id += 1
            orden   += 1

        for email in g["tecnico"]:
            contactos.append({
                "id":         next_id,
                "laboratorio": LAB_COMPARTIDO,
                "nombre":     email,
                "email":      email,
                "cargo":      "Técnico",
                "tipo":       "resultado_interno",
                "sold_to":    sold_to,
                "ship_to":    ship_to,
                "especie":    especie,
                "tipo_copia": "cc",
                "activo":     True,
                "orden":      orden,
            })
            next_id += 1
            orden   += 1

        for email in g["admin"]:
            contactos.append({
                "id":         next_id,
                "laboratorio": LAB_COMPARTIDO,
                "nombre":     email,
                "email":      email,
                "cargo":      "Admin",
                "tipo":       "resultado_interno",
                "sold_to":    sold_to,
                "ship_to":    ship_to,
                "especie":    especie,
                "tipo_copia": "bcc",
                "activo":     True,
                "orden":      orden,
            })
            next_id += 1
            orden   += 1

    return contactos


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main():
    aplicar = "--aplicar" in sys.argv
    limpiar = "--limpiar" in sys.argv  # borra contactos de resultado antes de reimportar

    # Buscar el Excel: primero argumento posicional, luego junto al script
    excel_arg = next((a for a in sys.argv[1:] if not a.startswith("-")), None)
    excel_path = excel_arg or str(EXCEL_PATH)

    if not os.path.exists(excel_path):
        sys.exit(
            f"No se encontró el Excel en:\n  {excel_path}\n\n"
            "Ponlo en backend/scripts/importar_contactos_resultado_excel.xlsx\n"
            "o pasa la ruta como argumento:\n"
            "  .venv\\Scripts\\python.exe scripts\\importar_contactos_resultado.py "
            "\"C:\\ruta\\al\\archivo.xlsx\""
        )

    print(f"Leyendo: {excel_path}")
    ws = leer_excel(excel_path)
    grupos = construir_grupos(ws)

    # Cargar JSON existente (via config_store para soportar R2)
    from app import config_store
    todos: list[dict] = config_store.leer("contactos_laboratorio.json", [])

    # Con --limpiar se eliminan los de resultado para reimportar desde cero
    if limpiar:
        existentes = [c for c in todos if c.get("tipo") not in ("resultado_cliente", "resultado_interno")]
        print(f"--limpiar: se eliminan {len(todos) - len(existentes)} contactos de resultado existentes.")
    else:
        existentes = todos

    max_id = max((c["id"] for c in existentes), default=0)

    # Detectar duplicados: grupos cuya clave ya existe en el JSON
    claves_existentes: set[tuple] = set()
    for c in existentes:
        if c.get("tipo") in ("resultado_cliente", "resultado_interno"):
            claves_existentes.add((
                c.get("sold_to", ""),
                c.get("ship_to", ""),
                c.get("especie", ""),
            ))

    grupos_nuevos = [
        g for g in grupos
        if (g["sold_to"], g["ship_to"], g["especie"]) not in claves_existentes
    ]
    grupos_duplicados = len(grupos) - len(grupos_nuevos)

    nuevos_contactos = grupos_a_contactos(grupos_nuevos, id_inicio=max_id + 1)

    # --- Resumen ---
    print()
    print("=" * 60)
    print("RESUMEN")
    print("=" * 60)
    print(f"Grupos en el Excel (sin Gonzalez):  {len(grupos)}")
    print(f"  Ya existían (se omiten):          {grupos_duplicados}")
    print(f"  Nuevos a importar:                {len(grupos_nuevos)}")
    print(f"Contactos a crear:                  {len(nuevos_contactos)}")
    print(f"ID actual máximo:                   {max_id}")
    print(f"Próximo ID:                         {max_id + 1}")
    print()

    if not grupos_nuevos:
        print("No hay grupos nuevos que importar.")
        return

    # Preview de primeros 10 grupos nuevos
    print("Primeros grupos nuevos:")
    for g in grupos_nuevos[:10]:
        esp = g["especie"] or "(global)"
        print(f"  {g['sold_to']} · {g['ship_to']} · {esp}")
        print(f"    clientes: {g['clientes']}")
        if g["comercial"]:
            print(f"    comercial (cc):  {g['comercial']}")
        if g["tecnico"]:
            print(f"    técnico (cc):    {g['tecnico']}")
        if g["admin"]:
            print(f"    admin (bcc):     {g['admin']}")

    if len(grupos_nuevos) > 10:
        print(f"  ... y {len(grupos_nuevos) - 10} grupos más.")

    print()

    if not aplicar:
        print(">> DRY RUN — no se escribió nada.")
        print(">> Agrega --aplicar para guardar.")
        return

    # Escribir (via config_store para soportar R2)
    resultado = existentes + nuevos_contactos
    config_store.escribir("contactos_laboratorio.json", resultado)

    destino = "R2" if __import__("app.r2", fromlist=["disponible"]).disponible() else JSON_PATH
    print(f"✓ Guardado en: {destino}")
    print(f"  Total contactos ahora: {len(resultado)}")


if __name__ == "__main__":
    main()
