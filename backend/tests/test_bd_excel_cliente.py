"""«Descargar BD» para una cuenta de cliente: el mismo formato, solo de SU Ship To y sin las columnas internas."""
from contextlib import contextmanager
import asyncio
import io

import openpyxl
import pytest
from fastapi import HTTPException

from app import reportes
from app.auth import Usuario


def _usuario(tipo="cliente", cliente="DOLE CHILE S.A.", planta="DOLE SAN FERNANDO"):
    return Usuario(id="1", email="x@dole.cl", nombre="Dole", tipoAcceso=tipo, clienteNombre=cliente, plantaNombre=planta)


def _fila():
    return {
        "id": 1, "nro_informe": "2026-1880-PC", "nro_solicitud": "OT-QUI0001", "laboratorio": "QUITECA", "sold_to": "DOLE CHILE S.A.",
        "ship_to": "DOLE SAN FERNANDO", "resultados": {"FDL": {"valor": 1.2, "nombre": "Fludioxonil"}}, "dosis": {}, "tipo_aplicacion": None,
        "email_solicitante": "interno@agrofresh.com", "email_laboratorio": "lab@quiteca.cl", "generado_por": "Jorge",
        "lista_para": "a@x.cl", "lista_cc": "b@x.cl", "lista_cco": "oculto@agrofresh.com",
    }


@pytest.fixture
def llamar(monkeypatch):
    capturado = {}

    @contextmanager
    def conexion(escribir=True):
        yield object()

    @contextmanager
    def cursor_dict(_c):
        yield object()

    def filas_de_bd(cur, filtro, params, ids, correos):
        capturado.update(filtro=filtro, params=params, ids=ids, correos=correos)
        return [_fila()]

    monkeypatch.setattr(reportes, "conexion", conexion)
    monkeypatch.setattr(reportes, "cursor_dict", cursor_dict)
    monkeypatch.setattr(reportes, "filas_de_bd", filas_de_bd)

    def _llamar(usuario, ids=None):
        resp = reportes.bd_excel(reportes.BdExcelIn(solicitud_ids=ids), usuario)
        async def leer():
            return b"".join([c async for c in resp.body_iterator])

        cuerpo = asyncio.run(leer())
        wb = openpyxl.load_workbook(io.BytesIO(cuerpo))
        ws = wb["BD"]
        cols = [c.value for c in ws[2]]
        return capturado, dict(zip(cols, [c.value for c in ws[3]]))

    return _llamar


def test_el_cliente_baja_el_mismo_formato_pero_solo_de_su_ship_to(llamar):
    cap, fila = llamar(_usuario())
    assert cap["params"]["cliente"] == "DOLE CHILE S.A." and cap["params"]["planta"] == "DOLE SAN FERNANDO"
    assert fila["N° Informe"] == "2026-1880-PC" and fila["Ship To"] == "DOLE SAN FERNANDO" and fila["FDL"] == 1.2


def test_el_cliente_no_ve_correos_ni_listas_de_distribucion(llamar):
    cap, fila = llamar(_usuario())
    assert cap["correos"] is None
    for columna in ("Email Solicitante", "Email Laboratorio", "Lista de Distribución (Para)", "Lista de Distribución (CC)",
                    "Lista de Distribución (CCO)", "Generado Por"):
        assert fila[columna] is None, columna


def test_el_personal_interno_sigue_viendo_todo_sin_acotar(llamar):
    cap, fila = llamar(_usuario(tipo="admin_general", cliente=None, planta=None))
    assert "cliente" not in cap["params"] and cap["correos"] is not None
    assert fila["Lista de Distribución (CCO)"] == "oculto@agrofresh.com" and fila["Email Solicitante"] == "interno@agrofresh.com"


def test_un_cliente_sin_cliente_ni_sucursal_no_baja_nada(llamar):
    with pytest.raises(HTTPException) as e:
        llamar(_usuario(cliente=None, planta=None))
    assert e.value.status_code == 403
