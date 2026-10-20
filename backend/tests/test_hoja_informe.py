"""
Hoja «Para el informe» del Excel de Diagnofruit: la cadena de la solicitud en posiciones
fijas, Sold To // Ship To y la lista de distribución. Solo la llevan las solicitudes de
Diagnofruit creadas con la marca `hoja_informe_v1`.
"""
import io

from openpyxl import load_workbook

from app import hoja_informe
from app.solicitud_excel import construir_workbook

DFN = dict(
    laboratorio="DIAGNOFRUIT", numero_solicitud="OT-DIAG0007", fecha_solicitud="2026-10-01",
    sold_to="COMERCIALIZADORA GARATE HERMANOS LI", ship_to="GREENEX PLANTA CODEGUA",
    nombre_muestreador="Barly Fuentes", email_solicitante="bfuentes@agrofresh.com",
    posicion_muestreo="Pozo vaciado", fecha_muestreo="2026-10-01", hora_muestreo="13:07",
    lote="2036", csg_productor="95848", csg_packing=None, numero_camara=None, numero_orden=None,
    kilos_procesados=None, producto_utilizado=None, linea_proceso="1", especie="Manzana", variedad="Fuji",
    tipo_muestra="Agua", observacion="sin obs", campos_laboratorio={"Tipo Aplicación": "Línea de proceso"},
    analitos_solicitados=["BOT"], hoja_informe_v1=True,
)
LISTAS = {"para": ["a@cliente.cl", "b@cliente.cl"], "cc": ["cc@x.cl"], "bcc": []}


def _libro(datos, listas=LISTAS):
    wb = construir_workbook(datos, [])
    return wb


def test_cadena_en_posicion_fija_con_vacios_entre_separadores():
    cadena = hoja_informe.cadena_solicitud(DFN)
    partes = hoja_informe.partes_de_cadena(cadena)
    assert len(partes) == len(hoja_informe.CAMPOS_CADENA)
    assert partes[0] == "OT-DIAG0007"
    assert partes[1] == "Línea de proceso"
    assert partes[2:5] == ["01-10-2026", "01-10-2026", "13:07"]
    # CSG packing, N° cámara, N° orden, kilos y producto vacíos: conservan su lugar
    assert cadena == (
        "OT-DIAG0007 // Línea de proceso // 01-10-2026 // 01-10-2026 // 13:07 // Manzana // Fuji // 1 "
        "// 95848 // // 2036 // Pozo vaciado // // // // // Agua // Barly Fuentes // bfuentes@agrofresh.com // // sin obs"
    )


def test_un_doble_slash_dentro_de_un_dato_no_corre_las_posiciones():
    d = {**DFN, "observacion": "a // b"}
    assert len(hoja_informe.partes_de_cadena(hoja_informe.cadena_solicitud(d))) == len(hoja_informe.CAMPOS_CADENA)


def test_gasto_sale_de_los_campos_del_servicio():
    d = {**DFN, "campos_laboratorio": {"Tipo Aplicación": "Actimist", "Gasto": "12"}}
    partes = hoja_informe.partes_de_cadena(hoja_informe.cadena_solicitud(d))
    nombres = [n for _, n in hoja_informe.CAMPOS_CADENA]
    assert partes[nombres.index("Tipo de servicio")] == "Actimist"
    assert partes[nombres.index("Gasto")] == "12"


def test_hoja_con_las_tres_cosas(monkeypatch):
    monkeypatch.setattr(hoja_informe, "destinatarios", lambda datos: LISTAS)
    wb = construir_workbook(DFN, [])
    assert hoja_informe.NOMBRE_HOJA in wb.sheetnames
    assert wb.sheetnames[-1] == "_data"
    ws = wb[hoja_informe.NOMBRE_HOJA]
    assert ws["B4"].value == hoja_informe.cadena_solicitud(DFN)
    assert ws["B5"].value == "COMERCIALIZADORA GARATE HERMANOS LI // GREENEX PLANTA CODEGUA"
    assert ws["B6"].value == "a@cliente.cl; b@cliente.cl"
    assert ws["B7"].value == "cc@x.cl"
    assert ws["B8"].value is None or ws["B8"].value == ""
    assert ws["C3"].value == "Ubicación"
    assert ws["C5"].value == "Predio/planta"
    assert [ws[f"C{n}"].value for n in (4, 6, 7, 8)] == ["OBSERVACIONES ADICIONALES"] * 4


def test_el_excel_se_guarda_y_se_vuelve_a_leer(monkeypatch):
    monkeypatch.setattr(hoja_informe, "destinatarios", lambda datos: LISTAS)
    buf = io.BytesIO()
    construir_workbook(DFN, []).save(buf)
    buf.seek(0)
    assert hoja_informe.NOMBRE_HOJA in load_workbook(buf).sheetnames


def test_solicitud_antigua_sin_marca_no_cambia():
    d = {k: v for k, v in DFN.items() if k != "hoja_informe_v1"}
    assert hoja_informe.NOMBRE_HOJA not in construir_workbook(d, []).sheetnames


def test_otros_laboratorios_no_llevan_la_hoja():
    for lab in ("ALS", "QUITECA", "AGROFRESH"):
        d = {**DFN, "laboratorio": lab}
        assert hoja_informe.NOMBRE_HOJA not in construir_workbook(d, []).sheetnames


def test_marcar_solo_diagnofruit():
    d = {"laboratorio": "Diagnofruit"}
    hoja_informe.marcar(d)
    assert d[hoja_informe.MARCA_HOJA_INFORME] is True
    o = {"laboratorio": "ALS"}
    hoja_informe.marcar(o)
    assert hoja_informe.MARCA_HOJA_INFORME not in o
