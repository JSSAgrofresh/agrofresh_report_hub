"""
Descarga de la BD de resultados (Report → Laboratorio).

Lo que importa aquí: que las columnas se acoten a lo que hay en la descarga
(un laboratorio filtrado no arrastra los analitos de los otros), que el
resultado aparezca donde en Solicitudes iba el ✓, y que un archivo filtrado
avise que no es la base completa.

La primera parte es pura. La segunda necesita Postgres con el esquema y se
salta sola si no lo hay.
"""
import io

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app import bd_excel, config_store
from app.auth import Usuario, usuario_actual
from app.main import app
from app.toma_muestras import ANALITOS_DEFECTO
from tests.utiles_bd import hay_base


def _fila(**kw):
    base = {"nro_informe": "INF-1", "laboratorio": "QUITECA", "resultados": {}, "dosis": {}, "tipo_aplicacion": None}
    base.update(kw)
    return base


def _encabezados(wb):
    ws = wb["BD"]
    return [c.value for c in ws[2]], [c.value for c in ws[1] if c.value]


FUNGICIDA = _fila(
    resultados={"FDL": {"valor": 0.42, "nombre": "Fludioxonil"}, "IMZ": {"valor": 1.5, "nombre": "Imazalil"}},
    dosis={"FDL": 5, "IMZ": 3},
    tipo_aplicacion="Línea de proceso",
)
AGUA_ALS = _fila(
    laboratorio="ALS",
    nro_informe="INF-2",
    resultados={"ECOLI": {"valor": 10, "nombre": "E. Coli"}, "COLT": {"valor": "<1", "nombre": "Coliformes"}},
)


def test_solo_salen_los_grupos_y_analitos_que_tienen_datos():
    wb = bd_excel.construir_workbook_bd([AGUA_ALS], ANALITOS_DEFECTO)
    columnas, bandas = _encabezados(wb)
    assert "E. Coli UFC/100mL" in columnas and "Coliformes Totales UFC/100mL" in columnas
    assert "FDL" not in columnas and "Levaduras UFC/mL" not in columnas
    assert bandas == ["GENERAL", "ALS — MICROBIOLOGÍA AGUA (FSMA)"]


def test_el_resultado_va_donde_en_solicitudes_iba_el_check_y_la_dosis_al_lado():
    wb = bd_excel.construir_workbook_bd([FUNGICIDA], ANALITOS_DEFECTO)
    columnas, _ = _encabezados(wb)
    ws = wb["BD"]
    fila = dict(zip(columnas, [c.value for c in ws[3]]))
    assert fila["FDL"] == 0.42 and fila["FDL Dosis"] == 5
    assert fila["IMZ"] == 1.5
    assert fila["Tipo Aplicación"] == "Línea de proceso"
    assert "✓" not in fila.values()
    # FDL y su dosis quedan pegadas, como en el formato de Solicitudes
    assert columnas.index("FDL Dosis") == columnas.index("FDL") + 1


def test_un_texto_como_menor_que_se_conserva():
    wb = bd_excel.construir_workbook_bd([AGUA_ALS], ANALITOS_DEFECTO)
    columnas, _ = _encabezados(wb)
    fila = dict(zip(columnas, [c.value for c in wb["BD"][3]]))
    assert fila["Coliformes Totales UFC/100mL"] == "<1"


def test_filtrar_por_ingrediente_deja_solo_ese_y_su_dosis():
    wb = bd_excel.construir_workbook_bd([FUNGICIDA], ANALITOS_DEFECTO, {"IMZ"})
    columnas, _ = _encabezados(wb)
    assert "IMZ" in columnas and "IMZ Dosis" in columnas
    assert "FDL" not in columnas and "FDL Dosis" not in columnas


def test_un_analito_sin_columna_en_el_formato_no_se_pierde():
    fila = _fila(resultados={"XYZ": {"valor": 7, "nombre": "Nuevo analito"}})
    wb = bd_excel.construir_workbook_bd([fila], ANALITOS_DEFECTO)
    columnas, bandas = _encabezados(wb)
    assert "Nuevo analito (XYZ)" in columnas
    assert bd_excel.TITULO_OTROS in bandas


def test_sin_resultados_solo_quedan_las_columnas_generales():
    wb = bd_excel.construir_workbook_bd([_fila()], ANALITOS_DEFECTO)
    columnas, bandas = _encabezados(wb)
    assert bandas == ["GENERAL"]
    assert columnas[:2] == ["N° Informe", "N° Solicitud"]


def test_la_hoja_de_aviso_solo_existe_en_una_descarga_filtrada():
    completa = bd_excel.construir_workbook_bd([FUNGICIDA], ANALITOS_DEFECTO)
    assert completa.sheetnames == ["BD"]
    filtrada = bd_excel.construir_workbook_bd([FUNGICIDA], ANALITOS_DEFECTO, nota_filtro="Laboratorio: QUITECA")
    assert filtrada.sheetnames == ["BD", "Filtros aplicados"]
    assert "NO es la base de datos completa" in filtrada["Filtros aplicados"]["A1"].value


# ── Contra Postgres ──────────────────────────────────────────────────────

con_bd = pytest.mark.skipif(not hay_base("producto_aplicado"), reason="sin Postgres con el esquema")


@pytest.fixture
def datos_bd(monkeypatch):
    from app.db import conexion, cursor_dict

    monkeypatch.setattr(config_store, "leer", lambda nombre, defecto: ANALITOS_DEFECTO)
    ids = {}
    with conexion() as conn, cursor_dict(conn) as cur:
        cur.execute("INSERT INTO cliente (nombre) VALUES ('__CLIENTE_BD__') RETURNING id")
        cli = cur.fetchone()["id"]
        cur.execute("INSERT INTO planta (cliente_id, nombre) VALUES (%s, '__PLANTA_BD__') RETURNING id", (cli,))
        pl = cur.fetchone()["id"]
        cur.execute("INSERT INTO analito (codigo, nombre, laboratorio) VALUES ('FDL', 'Fludioxonil', '__LAB_BD__') RETURNING id")
        fdl = cur.fetchone()["id"]
        cur.execute("INSERT INTO analito (codigo, nombre, laboratorio) VALUES ('ECOLI', 'E. Coli', '__LAB_BD__') RETURNING id")
        eco = cur.fetchone()["id"]
        for clave, lab, nro in (("quiteca", "QUITECA", "__INF_Q__"), ("als", "ALS", "__INF_A__")):
            cur.execute(
                "INSERT INTO solicitud (nro_solicitud, laboratorio, fecha_muestreo, especie, planta_id, referencia)"
                " VALUES (%s, %s, '2026-09-01', 'Cereza', %s, %s) RETURNING id",
                (nro, lab, pl, "OT-" + clave),
            )
            ids[clave] = cur.fetchone()["id"]
        cur.execute("INSERT INTO resultado (solicitud_id, analito_id, valor_num) VALUES (%s, %s, 0.42)", (ids["quiteca"], fdl))
        cur.execute("INSERT INTO producto_aplicado (solicitud_id, analito_id, dosis, tipo_aplicacion) VALUES (%s, %s, 5, 'Fogger')", (ids["quiteca"], fdl))
        cur.execute("INSERT INTO resultado (solicitud_id, analito_id, valor_num) VALUES (%s, %s, 10)", (ids["als"], eco))
    yield ids
    with conexion() as conn, cursor_dict(conn) as cur:
        cur.execute("DELETE FROM solicitud WHERE nro_solicitud LIKE '\\_\\_INF\\_%'")
        cur.execute("DELETE FROM analito WHERE laboratorio = '__LAB_BD__'")
        cur.execute("DELETE FROM planta WHERE nombre = '__PLANTA_BD__'")
        cur.execute("DELETE FROM cliente WHERE nombre = '__CLIENTE_BD__'")


@pytest.fixture
def cliente_http():
    quien = Usuario(id="2", email="a@x.cl", nombre="A", tipoAcceso="analista", area="cromatografia")
    app.dependency_overrides[usuario_actual] = lambda: quien
    yield TestClient(app)
    app.dependency_overrides.pop(usuario_actual, None)


def _hoja(resp):
    wb = load_workbook(io.BytesIO(resp.content))
    ws = wb["BD"]
    return wb, [c.value for c in ws[2]], [[c.value for c in fila] for fila in ws.iter_rows(min_row=3)]


@con_bd
def test_descarga_completa_trae_todo_y_no_avisa(datos_bd, cliente_http):
    r = cliente_http.post("/api/reportes/bd/excel", json={})
    assert r.status_code == 200 and "completa" in r.headers["content-disposition"]
    wb, columnas, filas = _hoja(r)
    assert wb.sheetnames == ["BD"]
    informes = {f[columnas.index("N° Informe")] for f in filas}
    assert {"__INF_Q__", "__INF_A__"} <= informes
    assert "FDL" in columnas and "E. Coli UFC/100mL" in columnas


@con_bd
def test_descarga_filtrada_solo_trae_lo_pedido_y_avisa(datos_bd, cliente_http):
    r = cliente_http.post(
        "/api/reportes/bd/excel",
        json={"solicitud_ids": [datos_bd["als"]], "descripcion_filtros": "Laboratorio: ALS"},
    )
    assert "filtrada" in r.headers["content-disposition"]
    wb, columnas, filas = _hoja(r)
    assert [f[columnas.index("N° Informe")] for f in filas] == ["__INF_A__"]
    assert "E. Coli UFC/100mL" in columnas and "FDL" not in columnas  # las columnas de QUITECA no vienen
    assert wb.sheetnames == ["BD", "Filtros aplicados"]
    assert wb["Filtros aplicados"]["A2"].value == "Laboratorio: ALS"


@con_bd
def test_trae_resultado_y_dosis_de_la_base(datos_bd, cliente_http):
    r = cliente_http.post("/api/reportes/bd/excel", json={"solicitud_ids": [datos_bd["quiteca"]]})
    _, columnas, filas = _hoja(r)
    fila = dict(zip(columnas, filas[0]))
    assert float(fila["FDL"]) == pytest.approx(0.42)
    assert float(fila["FDL Dosis"]) == 5
    assert fila["Tipo Aplicación"] == "Fogger"
    assert fila["N° Solicitud"] == "OT-quiteca"


@con_bd
def test_una_lista_vacia_no_baja_todo(datos_bd, cliente_http):
    # Un filtro que no deja nada NO puede terminar descargando la base entera.
    r = cliente_http.post("/api/reportes/bd/excel", json={"solicitud_ids": []})
    _, _, filas = _hoja(r)
    assert filas == []


def test_un_texto_de_filtro_no_puede_ser_una_formula():
    from app.reportes import _texto_seguro_excel

    assert _texto_seguro_excel("=HYPERLINK(\"x\")").startswith("'")
    assert _texto_seguro_excel("Laboratorio: ALS") == "Laboratorio: ALS"
    assert _texto_seguro_excel("  ") is None


def test_cliente_no_puede_usar_la_descarga():
    quien = Usuario(id="9", email="c@x.cl", nombre="C", tipoAcceso="cliente", area="postventa", clienteNombre="X")
    app.dependency_overrides[usuario_actual] = lambda: quien
    try:
        assert TestClient(app).post("/api/reportes/bd/excel", json={}).status_code == 403
    finally:
        app.dependency_overrides.pop(usuario_actual, None)


def _banda_fungicidas(filas):
    _, bandas = _encabezados(bd_excel.construir_workbook_bd(filas, ANALITOS_DEFECTO))
    return next(b for b in bandas if "FUNGICIDAS" in b)


def test_la_banda_de_fungicidas_nombra_solo_el_laboratorio_que_hay():
    quiteca = _fila(laboratorio="QUITECA", resultados={"FDL": {"valor": 1, "nombre": "F"}})
    agro = _fila(laboratorio="Agrofresh", resultados={"IMZ": {"valor": 2, "nombre": "I"}})
    assert _banda_fungicidas([quiteca]) == "QUITECA — ANÁLISIS DE RESIDUOS DE FUNGICIDAS"
    assert _banda_fungicidas([agro]) == "AGROFRESH — ANÁLISIS DE RESIDUOS DE FUNGICIDAS"
    assert _banda_fungicidas([quiteca, agro]) == "QUITECA / AGROFRESH — ANÁLISIS DE RESIDUOS DE FUNGICIDAS"


def test_un_laboratorio_sin_resultados_de_fungicidas_no_entra_en_la_banda():
    quiteca = _fila(laboratorio="QUITECA", resultados={"FDL": {"valor": 1, "nombre": "F"}})
    agua_agro = _fila(laboratorio="AGROFRESH", resultados={"ECOLI": {"valor": 3, "nombre": "E"}})
    assert _banda_fungicidas([quiteca, agua_agro]) == "QUITECA — ANÁLISIS DE RESIDUOS DE FUNGICIDAS"


def test_laboratorio_combinado_o_desconocido_usa_el_rotulo_de_los_dos():
    combinada = _fila(laboratorio="Quiteca / AgroFresh", resultados={"FDL": {"valor": 1, "nombre": "F"}})
    otra = _fila(laboratorio="???", resultados={"FDL": {"valor": 1, "nombre": "F"}})
    assert _banda_fungicidas([combinada]) == "QUITECA / AGROFRESH — ANÁLISIS DE RESIDUOS DE FUNGICIDAS"
    assert _banda_fungicidas([otra]) == "QUITECA / AGROFRESH — ANÁLISIS DE RESIDUOS DE FUNGICIDAS"


# ── Lo que la base no trae se completa desde la solicitud ────────────────

SOLICITUD_OT = {
    "numero_solicitud": "OT-ALS0004",
    "linea_proceso": "1",
    "numero_camara": "3",
    "kilos_procesados": 1.0,
    "producto_utilizado": "FUNGAZIL",
    "tipo_muestra": "Agua",
    "nombre_muestreador": "Jorge Sandoval",
    "generado_por": "Jorge Sandoval",
    "email_solicitante": "jorge.sandoval@agrofresh.com",
    "posicion_muestreo": "Pozo Vaciado",
    "fecha_solicitud": "28-09-2026",
}


def test_se_llena_desde_la_solicitud_lo_que_la_base_no_trae():
    fila = _fila(nro_solicitud=None, fecha_muestreo=None)
    bd_excel.completar_fila(fila, SOLICITUD_OT, ["a@lab.cl", "b@lab.cl"])
    assert fila["nro_solicitud"] == "OT-ALS0004"
    assert fila["tipo_muestra"] == "Agua"
    assert fila["nombre_muestreador"] == "Jorge Sandoval"
    assert fila["generado_por"] == "Jorge Sandoval"
    assert fila["email_solicitante"] == "jorge.sandoval@agrofresh.com"
    assert fila["email_laboratorio"] == "a@lab.cl; b@lab.cl"
    assert str(fila["fecha_solicitud"]) == "2026-09-28"  # texto DD-MM-AAAA -> fecha


def test_lo_que_la_base_ya_trae_no_se_pisa():
    fila = _fila(tipo_muestra="Fruta", lote="L-9")
    bd_excel.completar_fila(fila, {**SOLICITUD_OT, "lote": "otro"}, None)
    assert fila["tipo_muestra"] == "Fruta" and fila["lote"] == "L-9"


def test_solicitante_es_siempre_agrofresh():
    fila = _fila(solicitante="Otra persona")
    bd_excel.completar_fila(fila, None, None)
    assert fila["solicitante"] == "AGROFRESH"


def test_temporada_es_el_anio_de_la_muestra_y_no_pisa_una_existente():
    import datetime as dt

    fila = _fila(fecha_muestreo=dt.date(2026, 9, 25))
    bd_excel.completar_fila(fila, None, None)
    assert fila["temporada"] == 2026
    fila = _fila(fecha_muestreo=dt.date(2026, 9, 25), temporada="2025")
    bd_excel.completar_fila(fila, None, None)
    assert fila["temporada"] == "2025"
    sin_fecha = _fila()
    bd_excel.completar_fila(sin_fecha, None, None)
    assert sin_fecha.get("temporada") is None


def test_lo_que_no_hay_en_ningun_lado_queda_vacio_sin_romper():
    fila = _fila()
    bd_excel.completar_fila(fila, None, None)
    assert fila.get("numero_camara") is None and fila.get("kilos_procesados") is None
    wb = bd_excel.construir_workbook_bd([fila], ANALITOS_DEFECTO)
    assert wb["BD"].max_row == 3


def test_n_orden_ya_no_es_una_columna():
    columnas, _ = _encabezados(bd_excel.construir_workbook_bd([_fila()], ANALITOS_DEFECTO))
    assert "N° Orden" not in columnas


@pytest.mark.skipif(not hay_base("solicitud_archivo"), reason="sin la tabla del índice de solicitudes (0020)")
def test_la_descarga_une_con_el_indice_de_solicitudes(datos_bd, cliente_http, monkeypatch):
    import json

    from app import reportes
    from app.db import conexion, cursor_dict

    monkeypatch.setattr(reportes, "_correos_de_laboratorio", lambda lab: ["lab1@x.cl", "lab2@x.cl"])
    with conexion() as conn, cursor_dict(conn) as cur:
        cur.execute(
            "INSERT INTO solicitud_archivo (archivo, numero_solicitud, datos) VALUES ('__bd_test.xlsx', 'OT-als', %s)",
            (json.dumps({**SOLICITUD_OT, "numero_solicitud": "OT-als"}),),
        )
    try:
        r = cliente_http.post("/api/reportes/bd/excel", json={"solicitud_ids": [datos_bd["als"]]})
        _, columnas, filas = _hoja(r)
        fila = dict(zip(columnas, filas[0]))
        assert fila["Tipo Muestra"] == "Agua"
        assert fila["Nombre Muestreador"] == "Jorge Sandoval"
        assert fila["Email Solicitante"] == "jorge.sandoval@agrofresh.com"
        assert fila["Email Laboratorio"] == "lab1@x.cl; lab2@x.cl"
        assert fila["Solicitante"] == "AGROFRESH"
        assert fila["Temporada"] == 2026
    finally:
        with conexion() as conn, cursor_dict(conn) as cur:
            cur.execute("DELETE FROM solicitud_archivo WHERE archivo = '__bd_test.xlsx'")


@pytest.mark.skipif(not hay_base("solicitud_archivo"), reason="sin la tabla del índice de solicitudes (0020)")
def test_la_bd_trae_el_cruce_el_gasto_y_la_lista_de_distribucion(datos_bd, cliente_http, monkeypatch):
    """Las mismas columnas que la base «con muestra»: N° Muestra, recepción,
    Gasto, ensayo y la lista de distribución de resultados."""
    import json

    from app import reportes, toma_muestras as tm
    from app.db import conexion, cursor_dict

    monkeypatch.setattr(reportes, "_correos_de_laboratorio", lambda lab: [])
    contactos = [
        {"email": "cliente@x.cl", "tipo": "resultado_cliente", "activo": True, "orden": 1,
         "sold_to": "__CLIENTE_BD__", "ship_to": "__PLANTA_BD__", "especie": ""},
        {"email": "tec@agrofresh.com", "tipo": "resultado_interno", "tipo_copia": "bcc", "activo": True, "orden": 2,
         "sold_to": "__CLIENTE_BD__", "ship_to": "__PLANTA_BD__", "especie": ""},
    ]
    leer = tm._leer_config
    monkeypatch.setattr(
        tm, "_leer_config",
        lambda archivo, defecto=None: contactos if archivo == "contactos_laboratorio.json" else leer(archivo, defecto),
    )
    datos = {**SOLICITUD_OT, "numero_solicitud": "OT-als", "campos_laboratorio": {"Gasto": "12 L", "Código de Ensayo": "E-9"}}
    with conexion() as conn, cursor_dict(conn) as cur:
        cur.execute(
            "INSERT INTO solicitud_archivo (archivo, numero_solicitud, datos, codigo_muestra, cruzado_en)"
            " VALUES ('__bd_test2.xlsx', 'OT-als', %s, 'AGF0007', '2026-09-02T18:03:00+00:00')",
            (json.dumps(datos),),
        )
        cur.execute("UPDATE solicitud SET codigo_ensayo = NULL WHERE id = %s", (datos_bd["als"],))
    try:
        # El Excel de ALS no tiene fungicidas: se pide el de Quiteca, que sí.
        r = cliente_http.post("/api/reportes/bd/excel", json={"solicitud_ids": [datos_bd["quiteca"], datos_bd["als"]]})
        _, columnas, filas = _hoja(r)
        assert [c for c in columnas if c.startswith("Lista de Distribución")] == [
            "Lista de Distribución (Para)", "Lista de Distribución (CC)", "Lista de Distribución (CCO)",
        ]
        als = dict(zip(columnas, next(f for f in filas if f[columnas.index("N° Informe")] == "__INF_A__")))
        assert als["N° Muestra"] == "AGF0007"
        assert als["Fecha Recepción"] is not None and als["Hora Recepción"]
        assert als["Lista de Distribución (Para)"] == "cliente@x.cl"
        assert als["Lista de Distribución (CCO)"] == "tec@agrofresh.com"
        quiteca = dict(zip(columnas, next(f for f in filas if f[columnas.index("N° Informe")] == "__INF_Q__")))
        assert {"Gasto", "Código de Ensayo", "N° Ensayo"} <= set(columnas)
        assert quiteca["Lista de Distribución (Para)"] == "cliente@x.cl"
    finally:
        with conexion() as conn, cursor_dict(conn) as cur:
            cur.execute("DELETE FROM solicitud_archivo WHERE archivo = '__bd_test2.xlsx'")


# ── Enlace por parecido, cuando el resultado llegó sin OT ────────────────

def _cand(ot, lab="QUITECA", ship="DOLE PLANTA CODEGUA", esp="Palta", fecha="2026-09-23", **datos):
    import datetime as dt

    return {
        "numero_solicitud": ot, "laboratorio": lab, "ship_to": ship, "especie": esp,
        "fecha_muestreo": dt.date.fromisoformat(fecha), "datos": {"numero_solicitud": ot, **datos},
    }


def _fila_sin_ot(**kw):
    import datetime as dt

    base = {"laboratorio": "Quiteca", "ship_to": "DOLE PLANTA CODEGUA", "especie": "Palta", "fecha_muestreo": dt.date(2026, 9, 23)}
    base.update(kw)
    return base


def test_una_coincidencia_unica_enlaza_aunque_cambien_mayusculas_y_espacios():
    fila = _fila_sin_ot(ship_to="Dole  Planta Codegua", laboratorio="QUITECA")
    datos = bd_excel.buscar_por_parecido(fila, [_cand("OT-QUI0025", tipo_muestra="Fruta")])
    assert datos["numero_solicitud"] == "OT-QUI0025"


def test_dos_solicitudes_iguales_el_mismo_dia_no_se_adivinan():
    candidatas = [_cand("OT-QUI0019"), _cand("OT-QUI0020")]
    assert bd_excel.buscar_por_parecido(_fila_sin_ot(), candidatas) is None


def test_no_enlaza_si_cambia_la_fecha_el_laboratorio_la_planta_o_la_especie():
    fila = _fila_sin_ot()
    for otra in (
        _cand("OT-1", fecha="2026-09-22"),
        _cand("OT-2", lab="ALS"),
        _cand("OT-3", ship="OTRA PLANTA"),
        _cand("OT-4", esp="Limón"),
    ):
        assert bd_excel.buscar_por_parecido(fila, [otra]) is None


def test_sin_fecha_planta_o_especie_no_se_intenta():
    assert bd_excel.buscar_por_parecido(_fila_sin_ot(fecha_muestreo=None), [_cand("OT-1")]) is None
    assert bd_excel.buscar_por_parecido(_fila_sin_ot(ship_to=None), [_cand("OT-1")]) is None
    assert bd_excel.buscar_por_parecido(_fila_sin_ot(especie=""), [_cand("OT-1")]) is None


def test_dos_archivos_de_la_misma_solicitud_cuentan_como_una():
    # reindexar o un reanálisis puede dejar el mismo N° dos veces: no es ambigüedad
    assert bd_excel.buscar_por_parecido(_fila_sin_ot(), [_cand("OT-QUI0025"), _cand("OT-QUI0025")]) is not None


@pytest.mark.skipif(not hay_base("solicitud_archivo"), reason="sin la tabla del índice de solicitudes (0020)")
def test_la_descarga_enlaza_por_parecido_cuando_el_resultado_no_trae_ot(datos_bd, cliente_http, monkeypatch):
    import json

    from app import reportes
    from app.db import conexion, cursor_dict

    monkeypatch.setattr(reportes, "_correos_de_laboratorio", lambda lab: [])
    with conexion() as conn, cursor_dict(conn) as cur:
        # la solicitud del resultado de QUITECA: misma planta, especie y fecha, pero sin OT en la base
        cur.execute("SELECT ship_to_raw, planta_id FROM solicitud WHERE id = %s", (datos_bd["quiteca"],))
        cur.execute("SELECT nombre FROM planta WHERE nombre = '__PLANTA_BD__'")
        planta = cur.fetchone()["nombre"]
        cur.execute("UPDATE solicitud SET referencia = NULL WHERE id = %s", (datos_bd["quiteca"],))
        cur.execute(
            "INSERT INTO solicitud_archivo (archivo, numero_solicitud, laboratorio, ship_to, especie, fecha_muestreo, datos)"
            " VALUES ('__bd_par.xlsx', 'OT-QUI9999', 'QUITECA', %s, 'Cereza', '2026-09-01', %s)",
            (planta, json.dumps({"numero_solicitud": "OT-QUI9999", "tipo_muestra": "Fruta", "nombre_muestreador": "Catherine"})),
        )
    try:
        r = cliente_http.post("/api/reportes/bd/excel", json={"solicitud_ids": [datos_bd["quiteca"]]})
        _, columnas, filas = _hoja(r)
        fila = dict(zip(columnas, filas[0]))
        assert fila["N° Solicitud"] == "OT-QUI9999"
        assert fila["Nombre Muestreador"] == "Catherine"
        # con dos coincidencias ya no se adivina
        with conexion() as conn, cursor_dict(conn) as cur:
            cur.execute(
                "INSERT INTO solicitud_archivo (archivo, numero_solicitud, laboratorio, ship_to, especie, fecha_muestreo, datos)"
                " VALUES ('__bd_par2.xlsx', 'OT-QUI9998', 'QUITECA', %s, 'Cereza', '2026-09-01', %s)",
                (planta, json.dumps({"numero_solicitud": "OT-QUI9998", "nombre_muestreador": "Otra"})),
            )
        r = cliente_http.post("/api/reportes/bd/excel", json={"solicitud_ids": [datos_bd["quiteca"]]})
        _, columnas, filas = _hoja(r)
        assert dict(zip(columnas, filas[0]))["Nombre Muestreador"] is None
    finally:
        with conexion() as conn, cursor_dict(conn) as cur:
            cur.execute("DELETE FROM solicitud_archivo WHERE archivo LIKE '\\_\\_bd\\_par%'")
