"""Listas de distribución desde Administración General: comparar un Excel con
lo que hay y aplicar SOLO lo confirmado."""
import io

import openpyxl
import pytest

from app import listas_distribucion as ld

ADM = ["jorge.sandoval@agrofresh.com", "cguerrero@agrofresh.com", "agrofreshreporthub@gmail.com"]


def _c(id_, tipo, email, especie="", cargo="", copia="cc", activo=True, orden=1, sold="CLI SA", ship="PLANTA UNO"):
    return {"id": id_, "tipo": tipo, "email": email, "especie": especie, "cargo": cargo, "tipo_copia": copia,
            "activo": activo, "orden": orden, "sold_to": sold, "ship_to": ship, "laboratorio": "AGROFRESH", "nombre": email}


def _sistema():
    """Planta con lista general de cliente, un comercial, un técnico MAL puesto (cc) y dos admin."""
    return [
        _c(1, "resultado_cliente", "cli@x.cl"),
        _c(2, "resultado_interno", "com@agrofresh.com", cargo="Comercial", copia="cc"),
        _c(3, "resultado_interno", "tec@agrofresh.com", cargo="Técnico", copia="cc"),
        _c(4, "resultado_interno", ADM[0], cargo="Admin", copia="bcc"),
        _c(5, "resultado_interno", ADM[1], cargo="Admin", copia="bcc"),
    ]


def _fila(**kw):
    base = {"sold_to": "CLI SA", "ship_to": "PLANTA UNO", "admin": [], "comercial": [], "tecnico": [],
            "clientes": {c: [] for c in ld.CATEGORIAS}}
    base.update(kw)
    return base


def _cambios(contactos, fila, listados=None):
    return ld.comparar(ld.estado_desde_contactos(contactos), [fila], listados)["cambios"]


def test_estado_expande_la_lista_general_a_todas_las_categorias_y_marca_copia_mal():
    e = ld.estado_desde_contactos(_sistema())[ld.clave_planta("CLI SA", "PLANTA UNO")]
    assert all(e["clientes"][c] == ["cli@x.cl"] for c in ld.CATEGORIAS)
    assert e["tecnico"] == ["tec@agrofresh.com"] and e["copia_mal"] == ["tec@agrofresh.com"]
    assert e["admin"] == ADM[:2]


def test_celda_vacia_no_quita_a_nadie_y_sin_diferencias_no_hay_cambios():
    fila = _fila(comercial=["com@agrofresh.com"], tecnico=["tec@agrofresh.com"], admin=ADM[:2],
                 clientes={c: ["cli@x.cl"] for c in ld.CATEGORIAS})
    assert [c["tipo"] for c in _cambios(_sistema(), fila)] == ["copia"]  # solo falta ajustar el técnico
    # una fila totalmente vacía tampoco quita a nadie: lo único que queda es el ajuste de copia
    assert [c["tipo"] for c in _cambios(_sistema(), _fila())] == ["copia"]


def test_agregar_el_correo_del_sistema_y_quitar_a_alguien():
    fila = _fila(admin=ADM, comercial=["otro@agrofresh.com"])
    por_campo = {c["campo"]: c for c in _cambios(_sistema(), fila)}
    assert por_campo["admin"]["agregar"] == [ADM[2]] and por_campo["admin"]["quitar"] == []
    assert por_campo["comercial"]["agregar"] == ["otro@agrofresh.com"]
    assert por_campo["comercial"]["quitar"] == ["com@agrofresh.com"]


def test_aplicar_solo_lo_confirmado():
    fila = _fila(admin=ADM, comercial=["otro@agrofresh.com"])
    cambios = _cambios(_sistema(), fila)
    solo_admin = [c for c in cambios if c["campo"] == "admin"]
    nuevos, r = ld.aplicar(_sistema(), solo_admin)
    assert r["aplicados"] == 1 and r["plantas"] == 1
    emails = {(c["email"], c["tipo_copia"]) for c in nuevos if c["tipo"] == "resultado_interno" and c["cargo"] == "Admin"}
    assert (ADM[2], "bcc") in emails
    assert any(c["email"] == "com@agrofresh.com" for c in nuevos)  # el comercial NO se tocó


def test_ajustar_copia_pasa_el_tecnico_a_copia_oculta():
    cambios = [c for c in _cambios(_sistema(), _fila(tecnico=["tec@agrofresh.com"])) if c["tipo"] == "copia"]
    nuevos, _ = ld.aplicar(_sistema(), cambios)
    assert [c["tipo_copia"] for c in nuevos if c["email"] == "tec@agrofresh.com"] == ["bcc"]


def test_agregar_un_cliente_a_una_categoria_separa_por_especie_sin_perder_a_nadie():
    fila = _fila(clientes={**{c: ["cli@x.cl"] for c in ld.CATEGORIAS}, "Kiwi": ["cli@x.cl", "nuevo@x.cl"]})
    cambios = _cambios(_sistema(), fila)
    assert [(c["campo"], c["agregar"]) for c in cambios if c["tipo"] == "campo"] == [("Kiwi", ["nuevo@x.cl"])]
    nuevos, _ = ld.aplicar(_sistema(), [c for c in cambios if c["tipo"] == "campo"])
    cli = [c for c in nuevos if c["tipo"] == "resultado_cliente"]
    assert not any(c["especie"] == "" for c in cli)  # ya no hay lista general
    assert {c["email"] for c in cli if c["especie"] == "Kiwi"} == {"cli@x.cl", "nuevo@x.cl"}
    assert {c["email"] for c in cli if c["especie"] == "Manzana"} == {"cli@x.cl"}
    assert len({c["id"] for c in nuevos}) == len(nuevos)  # ids únicos


def test_agregar_lo_mismo_a_todas_las_categorias_vuelve_a_la_lista_general():
    fila = _fila(clientes={c: ["cli@x.cl", "nuevo@x.cl"] for c in ld.CATEGORIAS})
    cambios = [c for c in _cambios(_sistema(), fila) if c["tipo"] == "campo"]
    assert len(cambios) == len(ld.CATEGORIAS)
    nuevos, _ = ld.aplicar(_sistema(), cambios)
    cli = [c for c in nuevos if c["tipo"] == "resultado_cliente"]
    assert {c["especie"] for c in cli} == {""} and {c["email"] for c in cli} == {"cli@x.cl", "nuevo@x.cl"}


def test_planta_nueva_se_crea_con_sus_roles_y_avisa_si_no_esta_en_listados():
    fila = _fila(sold_to="OTRO SA", ship_to="PLANTA DOS", comercial=["c@agrofresh.com"], tecnico=["t@agrofresh.com"],
                 admin=ADM[:1], clientes={"Kiwi": ["k@x.cl"]})
    cambios = _cambios(_sistema(), fila, listados={})
    assert cambios[0]["tipo"] == "planta_nueva" and "Listados" in cambios[0]["aviso"]
    nuevos, _ = ld.aplicar(_sistema(), cambios)
    copia = {c["email"]: c["tipo_copia"] for c in nuevos if c["ship_to"] == "PLANTA DOS" and c["tipo"] == "resultado_interno"}
    assert copia == {"c@agrofresh.com": "cc", "t@agrofresh.com": "bcc", ADM[0]: "bcc"}
    assert {c["especie"] for c in nuevos if c["ship_to"] == "PLANTA DOS" and c["tipo"] == "resultado_cliente"} == {"Kiwi"}


def test_el_nombre_de_listados_manda_al_crear():
    fila = _fila(sold_to="otro  sa", ship_to="planta  dos", comercial=["c@agrofresh.com"])
    lis = {ld.clave_planta("otro sa", "planta dos"): ("OTRO SA", "PLANTA DOS")}
    assert _cambios(_sistema(), fila, lis)[0]["planta"] == {"sold_to": "OTRO SA", "ship_to": "PLANTA DOS"}


def test_exportar_y_volver_a_subir_no_propone_cambios():
    ok = _sistema()
    ok[2]["tipo_copia"] = "bcc"  # técnico bien puesto
    estado = ld.estado_desde_contactos(ok)
    filas, avisos = ld.leer_filas_excel(ld.construir_excel(estado))
    assert avisos == []
    resultado = ld.comparar(estado, filas)
    assert resultado["cambios"] == [] and resultado["resumen"]["plantas_sin_cambios"] == 1


def test_lee_la_hoja_del_maestro_y_avisa_de_textos_que_no_son_correo():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = ld.HOJA_MAESTRO
    ws.append(["#", "ñ", "VIGENTE POST VENTA", "S NUMBER", "SOLD TO NAME", "SHIP TO NUMBER", "SHIP TO NAME",
               "Admin Report Hub", "Comercial a cargo", "Técnico a cargo", "Manzana y Pera", "Kiwi", "Carozos",
               "Cerezas", "Citricos", "Arandanos ", "Paltas", "Nueces y Pasas", "Granada", "FSIN CORREO"])
    ws.append([1, 1, "SI", 1, "CLI SA", 1, "PLANTA UNO", "a@x.cl;cguerrero", "#N/A", "NA", "m@x.cl", None, None,
               None, None, None, None, None, None, "OK"])
    ws.append([2, 2, "NO", 2, "CLI SA", 2, "PLANTA NO VIGENTE", None, None, None, None, None, None,
               None, None, None, None, None, None, "OK"])
    ws.append([3, 3, None, 3, "CLI SA", 3, "PLANTA SIN MARCA", None, None, None, None, None, None,
               None, None, None, None, None, None, "OK"])  # vigente vacío: cuenta, solo "NO" la saca
    buf = io.BytesIO()
    wb.save(buf)
    filas, avisos = ld.leer_filas_excel(buf.getvalue())
    assert len(filas) == 2 and filas[0]["admin"] == ["a@x.cl"] and filas[0]["comercial"] == []
    assert filas[0]["clientes"]["Manzana y Pera"] == ["m@x.cl"]
    assert any("cguerrero" in a for a in avisos)


def test_un_excel_sin_las_columnas_da_error_claro():
    wb = openpyxl.Workbook()
    wb.active.append(["a", "b"])
    buf = io.BytesIO()
    wb.save(buf)
    with pytest.raises(ValueError, match="SOLD TO NAME"):
        ld.leer_filas_excel(buf.getvalue())


def test_planta_que_no_esta_en_listados_sugiere_el_nombre_parecido():
    lis = {
        ld.clave_planta("EXPORTADORA AGUA SANTA SA", "AGUA SANTA PLANTA LISONJERAS"): ("EXPORTADORA AGUA SANTA SA", "AGUA SANTA PLANTA LISONJERAS"),
        ld.clave_planta("OTRO CLIENTE", "PLANTA SIN RELACION"): ("OTRO CLIENTE", "PLANTA SIN RELACION"),
    }
    fila = _fila(sold_to="EXPORTADORA AGUA SANTA S.A", ship_to="AGUA SANTA PLANTA LISONJERA", comercial=["c@agrofresh.com"])
    c = _cambios(_sistema(), fila, lis)[0]
    assert c["aviso"] and c["sugerencias"] == [
        {"sold_to": "EXPORTADORA AGUA SANTA SA", "ship_to": "AGUA SANTA PLANTA LISONJERAS"}]


# --- estado para la tabla ---------------------------------------------------

def test_estado_para_tabla_marca_copia_mal_por_rol_y_si_esta_en_listados():
    estado = ld.estado_desde_contactos(_sistema())
    clave = ld.clave_planta("CLI SA", "PLANTA UNO")
    r = ld.estado_para_tabla(estado, {clave: ("CLI SA", "PLANTA UNO"), ld.clave_planta("X", "Y"): ("X", "Y")})
    fila = r["filas"][0]
    assert fila["en_listados"] is True and fila["copia_mal"] == {"admin": [], "comercial": [], "tecnico": ["tec@agrofresh.com"]}
    assert r["resumen"] == {"plantas_con_lista": 1, "plantas_listados": 2, "listados_sin_lista": 1}
    assert len(r["filas"]) == 1  # la de Listados sin lista solo entra si se pide
    todas = ld.estado_para_tabla(estado, {clave: ("CLI SA", "PLANTA UNO"), ld.clave_planta("X", "Y"): ("X", "Y")}, True)
    assert [(f["sold_to"], f["sin_contactos"]) for f in todas["filas"]] == [("CLI SA", False), ("X", True)]  # orden alfabético


def test_estado_para_tabla_sin_base_no_inventa_nada():
    r = ld.estado_para_tabla(ld.estado_desde_contactos(_sistema()), None)
    assert r["filas"][0]["en_listados"] is None and r["resumen"]["plantas_listados"] is None


class _Cur:
    """Cursor de mentira: clientes y plantas ya existentes, y lo que se inserte."""

    def __init__(self, clientes, plantas):
        self.clientes, self.plantas, self.inserts, self._ultimo = clientes, plantas, [], None

    def execute(self, sql, params=()):
        self._ultimo = None
        if sql.startswith("SELECT id, nombre FROM cliente"):
            self._ultimo = list(self.clientes)
        elif sql.startswith("SELECT id, nombre FROM planta"):
            self._ultimo = [p for p in self.plantas if p["cliente_id"] == params[0]]
        elif sql.startswith("INSERT INTO cliente"):
            self.inserts.append(("cliente", params))
            self._ultimo = {"id": 99}
        elif sql.startswith("INSERT INTO planta"):
            self.inserts.append(("planta", params))
            self._ultimo = {"id": 100}

    def fetchall(self):
        return self._ultimo

    def fetchone(self):
        return self._ultimo


def test_crear_planta_en_listados_reusa_lo_que_existe_sin_duplicar():
    cur = _Cur([{"id": 1, "nombre": "EXPORTADORA AGUA SANTA SA"}],
               [{"id": 7, "cliente_id": 1, "nombre": "AGUA SANTA PLANTA LISONJERA"}])
    r = ld.asegurar_planta(cur, "exportadora  agua santa sa", "agua santa planta   lisonjera")
    assert cur.inserts == []
    assert r == {"sold_to": "EXPORTADORA AGUA SANTA SA", "ship_to": "AGUA SANTA PLANTA LISONJERA",
                 "cliente_creado": False, "planta_creada": False}


def test_crear_planta_en_listados_crea_cliente_y_planta_con_sus_codigos_sap():
    cur = _Cur([], [])
    r = ld.asegurar_planta(cur, "CLIENTE NUEVO SA", "PLANTA NUEVA", "10001", "20002")
    assert cur.inserts == [("cliente", ("CLIENTE NUEVO SA", "10001")), ("planta", (99, "PLANTA NUEVA", "20002"))]
    assert r["cliente_creado"] and r["planta_creada"]
    with pytest.raises(ValueError):
        ld.asegurar_planta(cur, "SOLO CLIENTE", "  ")


def test_el_excel_trae_los_codigos_sap_si_los_hay():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["S NUMBER", "SOLD TO NAME", "SHIP TO NUMBER", "SHIP TO NAME", "Admin Report Hub"])
    ws.append([10001, "CLI SA", 20002, "PLANTA UNO", "a@x.cl"])
    buf = io.BytesIO()
    wb.save(buf)
    filas, _ = ld.leer_filas_excel(buf.getvalue())
    assert (filas[0]["codigo_sold"], filas[0]["codigo_ship"]) == ("10001", "20002")


# --- plantas que el Excel nuevo ya no trae ---------------------------------

def _dos_plantas():
    otra = [_c(10 + i, c["tipo"], c["email"], cargo=c["cargo"], copia=c["tipo_copia"], sold="OTRO SA", ship="PLANTA DOS")
            for i, c in enumerate(_sistema())]
    return _sistema() + otra


def test_planta_que_no_viene_en_el_excel_se_ofrece_para_quitar_pero_no_se_quita_sola():
    contactos = _dos_plantas()
    r = ld.comparar(ld.estado_desde_contactos(contactos), [_fila(clientes={c: ["cli@x.cl"] for c in ld.CATEGORIAS})])
    assert r["retiradas"] == [{"planta": {"sold_to": "OTRO SA", "ship_to": "PLANTA DOS"}}]
    assert all(c["tipo"] != "planta_quitar" for c in r["cambios"])
    nuevos, _ = ld.aplicar(contactos, r["cambios"])      # sin confirmar nada, nadie se va
    assert {c["ship_to"] for c in nuevos} == {"PLANTA UNO", "PLANTA DOS"}


def test_confirmar_quitar_borra_toda_la_lista_de_esa_planta_y_solo_de_esa():
    contactos = _dos_plantas()
    cambio = {"id": "x", "tipo": "planta_quitar", "planta": {"sold_to": "OTRO SA", "ship_to": "PLANTA DOS"},
              "campo": "planta", "agregar": [], "quitar": [], "corregir": []}
    nuevos, hechos = ld.aplicar(contactos, [cambio])
    assert {c["ship_to"] for c in nuevos} == {"PLANTA UNO"} and len(nuevos) == len(_sistema())
    assert hechos["aplicados"] == 1 and hechos["plantas"] == 1
    # quitar algo que ya no está no falla
    _, otra = ld.aplicar(nuevos, [cambio])
    assert otra["aplicados"] == 0 and otra["ignorados"]


def test_quitar_una_planta_de_linea_de_proceso_no_toca_actimist():
    contactos = _dos_plantas() + [{**_c(99, "resultado_cliente", "act@x.cl", sold="OTRO SA", ship="PLANTA DOS"), "servicio": "actimist"}]
    cambio = {"id": "x", "tipo": "planta_quitar", "planta": {"sold_to": "OTRO SA", "ship_to": "PLANTA DOS"},
              "campo": "planta", "agregar": [], "quitar": [], "corregir": []}
    nuevos, _ = ld.aplicar(contactos, [cambio], "")
    assert [c["email"] for c in nuevos if c.get("servicio") == "actimist"] == ["act@x.cl"]


def test_excel_sin_filas_no_ofrece_quitar_nada():
    assert ld.comparar(ld.estado_desde_contactos(_dos_plantas()), [])["retiradas"] == []


def test_ryd_tiene_su_propia_lista_y_no_toca_la_de_linea():
    contactos = _sistema()
    fila = _fila(comercial=["com.ryd@agrofresh.com"], clientes={c: ["cli.ryd@x.cl"] for c in ld.CATEGORIAS})
    estado_ryd = ld.estado_desde_contactos(ld.del_servicio(contactos, "ryd"))
    assert estado_ryd == {}                                    # Línea de proceso no se cuela en RYD
    cambios = ld.comparar(estado_ryd, [fila], {ld.clave_planta("CLI SA", "PLANTA UNO"): ("CLI SA", "PLANTA UNO")})["cambios"]
    nuevos, _ = ld.aplicar(contactos, cambios, "ryd")
    ryd = [c for c in nuevos if c.get("servicio") == "ryd"]
    assert {c["email"] for c in ryd} == {"com.ryd@agrofresh.com", "cli.ryd@x.cl"}
    assert [c for c in nuevos if not c.get("servicio")] == contactos        # Línea intacta


def test_cada_lista_dice_a_quien_manda_siempre_aunque_no_tenga_plantas():
    from app.servicios import fijos_de_lista
    for lista, para in (("actimist", ["CJIMENEZ@AGROFRESH.COM", "CVALENZUELA@AGROFRESH.COM"]),
                        ("ecofog", ["CJIMENEZ@AGROFRESH.COM", "CVALENZUELA@AGROFRESH.COM"]),
                        ("ryd", ["CCACERES@AGROFRESH.COM", "FGONZALEZ@AGROFRESH.COM"])):
        f = ld._fijos_para_pantalla(lista)
        assert f["para"] == para == fijos_de_lista(lista)["para"]
        assert f["cc"] == ["JORGE.SANDOVAL@AGROFRESH.COM", "AGROFRESHREPORTHUB@GMAIL.COM"] and f["respaldo"] == []
    linea = ld._fijos_para_pantalla("")
    assert linea["para"] == [] and linea["cc"] == [] and "CGUERRERO@AGROFRESH.COM" in linea["respaldo"]


# --- cambiar el nombre de una planta sin perder su lista -------------------

def _lista_losonjera():
    mala = [_c(40 + i, c["tipo"], c["email"], cargo=c["cargo"], copia=c["tipo_copia"], sold="AGUA SANTA S.A", ship="PLANTA LOSONJERA")
            for i, c in enumerate(_sistema())]
    ryd = [{**c, "id": 60 + i, "servicio": "ryd"} for i, c in enumerate(mala)]
    actimist = [{**c, "id": 80 + i, "servicio": "actimist"} for i, c in enumerate(mala)]
    return _sistema() + mala + ryd + actimist


def test_renombrar_conserva_la_lista_y_alcanza_a_las_listas_que_comparten_el_listado():
    contactos = _lista_losonjera()
    nuevos, n = ld.renombrar_en_contactos(contactos, "AGUA SANTA S.A", "PLANTA LOSONJERA", "PLANTA LISONJERA", "")
    assert n == 2 * len(_sistema())                              # Línea de proceso + RYD (mismo listado)
    de = lambda lista, ship: sorted(c["email"] for c in nuevos if c.get("servicio", "") == lista and c["ship_to"] == ship)
    assert de("", "PLANTA LISONJERA") == sorted(c["email"] for c in _sistema())   # nadie se perdió
    assert de("ryd", "PLANTA LISONJERA") and not de("", "PLANTA LOSONJERA")
    assert de("actimist", "PLANTA LOSONJERA") and not de("actimist", "PLANTA LISONJERA")   # Actimist tiene su listado
    assert {c["ship_to"] for c in nuevos if c["ship_to"] == "PLANTA UNO"}                    # otras plantas intactas
    assert len(nuevos) == len(contactos)


def test_renombrar_a_un_nombre_que_ya_tiene_lista_se_rechaza():
    contactos = _lista_losonjera() + [_c(99, "resultado_cliente", "x@x.cl", sold="AGUA SANTA S.A", ship="PLANTA LISONJERA")]
    with pytest.raises(ValueError, match="Ya hay una planta"):
        ld.renombrar_en_contactos(contactos, "AGUA SANTA S.A", "PLANTA LOSONJERA", "PLANTA LISONJERA", "")


class _CurRenombrar:
    def __init__(self, clientes, plantas):
        self.clientes, self.plantas, self.ultimo, self.actualizado = clientes, plantas, [], None

    def execute(self, sql, params=None):
        if sql.startswith("SELECT id, nombre FROM cliente"):
            self.ultimo = self.clientes
        elif sql.startswith("SELECT id, nombre FROM planta"):
            self.ultimo = self.plantas
        elif sql.startswith("UPDATE planta"):
            self.actualizado = params

    def fetchall(self):
        return self.ultimo


def test_renombrar_en_listados_cambia_el_nombre_de_la_misma_planta():
    cur = _CurRenombrar([{"id": 1, "nombre": "AGUA SANTA S.A"}], [{"id": 7, "nombre": "PLANTA LOSONJERA"}, {"id": 8, "nombre": "PLANTA OTRA"}])
    r = ld.renombrar_planta_listados(cur, "", "agua santa s.a", "planta losonjera", "PLANTA  LISONJERA")
    assert cur.actualizado == ("PLANTA LISONJERA", 7) and r["de"] == "PLANTA LOSONJERA"
    with pytest.raises(ValueError, match="Ya existe"):
        ld.renombrar_planta_listados(cur, "", "AGUA SANTA S.A", "PLANTA LOSONJERA", "planta otra")
    with pytest.raises(ValueError, match="no está en Listados"):
        ld.renombrar_planta_listados(cur, "", "AGUA SANTA S.A", "PLANTA FANTASMA", "X")
