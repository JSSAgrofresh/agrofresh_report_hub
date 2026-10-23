"""El script de diagnóstico, con una base de mentira: arma el Excel y cuenta lo que debe."""
import importlib.util
import sys
from contextlib import contextmanager
from datetime import date
from pathlib import Path

from openpyxl import load_workbook

RUTA = Path(__file__).resolve().parents[1] / "scripts" / "diagnostico_informes_report.py"


def _cargar():
    spec = importlib.util.spec_from_file_location("diagnostico_informes_report", RUTA)
    mod = importlib.util.module_from_spec(spec)
    sys.modules["diagnostico_informes_report"] = mod
    spec.loader.exec_module(mod)
    return mod


class _Cursor:
    def __init__(self, tablas):
        self.tablas, self.ultimo = tablas, []

    def execute(self, sql, params=None):
        for clave, filas in self.tablas.items():
            if clave in sql:
                self.ultimo = filas
                return
        raise AssertionError(f"consulta inesperada: {sql[:60]}")

    def fetchall(self):
        return self.ultimo

    def __enter__(self):
        return self

    def __exit__(self, *a):
        return False


def test_el_script_arma_el_excel_y_encuentra_el_ecofog_mal_clasificado(tmp_path, monkeypatch, capsys):
    mod = _cargar()
    informes = [
        {"id": 1, "nro_solicitud": "2026-1908-PC", "laboratorio": "Quiteca", "referencia": "OT-QUI0049", "servicio": None, "tipo_servicio": "Linea de Proceso",
         "especie": "Manzana", "variedad": "Pink Lady", "planta_id": 5, "fecha_muestreo": date(2026, 10, 8), "fecha_entrada": date(2026, 10, 8),
         "fecha_recepcion": date(2026, 10, 8), "fecha_informe": date(2026, 10, 9), "fecha_analisis": date(2026, 10, 8), "sold_to": "AGRICOLA X", "ship_to": "PLANTA X"},
        {"id": 2, "nro_solicitud": "2026-1909-PC", "laboratorio": "Quiteca", "referencia": "OT-QUI0050", "servicio": None, "tipo_servicio": "Linea de Proceso",
         "especie": "Manzana", "variedad": None, "planta_id": 5, "fecha_muestreo": date(2026, 10, 8), "fecha_entrada": date(2026, 10, 8),
         "fecha_recepcion": date(2026, 10, 8), "fecha_informe": date(2026, 10, 9), "fecha_analisis": date(2026, 10, 8), "sold_to": "AGRICOLA X", "ship_to": "PLANTA X"},
    ]
    resultados = [{"solicitud_id": i, "analito_id": 1, "analito_raw": "FDL", "codigo": "FDL", "valor_num": 1.0, "valor_texto": None} for i in (1, 2)]
    base = {"laboratorio": "Quiteca", "sold_to": "AGRICOLA X", "ship_to": "PLANTA X", "especie": "Manzana", "fecha_muestreo": date(2026, 10, 8)}
    tomas = [
        {"ot": "OT-QUI0049", **base, "datos": {"campos_laboratorio": {"Tipo Aplicación": "Ecofog"}, "analitos_solicitados": ["FDL"]}},
        {"ot": "OT-QUI0050", **base, "datos": {"campos_laboratorio": {"Tipo Aplicación": "Línea de proceso"}, "analitos_solicitados": ["FDL"]}},
    ]
    tablas = {"FROM solicitud s": informes, "FROM resultado r": resultados, "FROM solicitud_archivo": tomas}

    @contextmanager
    def conexion(escribir=False):
        yield object()

    monkeypatch.setattr(mod, "conexion", conexion)
    cursor = _Cursor(tablas)
    monkeypatch.setattr(mod, "cursor_dict", lambda conn: cursor)
    salida = tmp_path / "d.xlsx"
    assert mod.main(["--salida", str(salida)]) == 0

    texto = capsys.readouterr().out
    assert "Informes revisados: 2" in texto and "sin hallazgos graves o medios: 1" in texto
    assert "SERVICIO_SIN_CLASIFICAR" in texto
    wb = load_workbook(salida)
    assert wb.sheetnames == ["Resumen", "Por laboratorio", "Hallazgos"]
    hallazgos = [[c.value for c in fila] for fila in wb["Hallazgos"].iter_rows(min_row=2)]
    fila = next(f for f in hallazgos if f[1] == "SERVICIO_SIN_CLASIFICAR")
    assert fila[0] == "alto" and fila[2] == "2026-1908-PC" and "Ecofog" in fila[10]
