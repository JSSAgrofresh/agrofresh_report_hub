"""Plan de entregables de los comerciales: lector de los Excel y cruce con las plantas del sistema."""
import io
from datetime import date

import openpyxl

from app.listas_distribucion import estado_desde_contactos
from app.plan_cruce import correos_de_comercial, plantas_de, proponer
from app.plan_entregables import categoria_servicio, leer_plan, leer_planes, lunes_de_semana


def _excel(filas, semanas=("W45", "W46", "W47"), total_pie=True) -> io.BytesIO:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["B2"] = "Plan Servicio de Analisis & Entrega de Información en Cerezas 2026-27  -  JC Labarca:"
    ws["B4"], ws["C4"] = "CLIENTE", "SERVICIO"
    for i, s in enumerate(semanas):
        ws.cell(4, 4 + i, s)
        ws.cell(5, 4 + i, "02-08 Nov")
    ws.cell(4, 4 + len(semanas), "TOTAL")
    r = 6
    for cliente, servicio, valores, total in filas:
        ws.cell(r, 2, cliente)
        ws.cell(r, 3, servicio)
        for i, v in enumerate(valores):
            if v is not None:
                ws.cell(r, 4 + i, v)
        ws.cell(r, 4 + len(semanas), total)
        r += 1
    if total_pie:
        ws.cell(r, 2, "TOTAL ")
        ws.cell(r + 1, 2, "Residuales")
        ws.cell(r + 1, 3, 999)
        ws.cell(r + 2, 2, "Simbologia Patologicos:")
        ws.cell(r + 3, 2, "1 = Corresponde a Pozo Fungicida.")
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


class TestLector:
    def test_lee_comercial_semanas_y_filas(self):
        p = leer_plan(_excel([("DOLE LONTUE", "Residuales L1", [1, 2, None], 3), ("DOLE LONTUE", "Patologias L2", [None, 4, 4], 8)]))
        assert p["comercial"] == "JC Labarca"
        assert p["semanas"] == ["W45", "W46", "W47"]
        assert [(f["servicio"], f["linea"], f["semana"], f["cantidad"]) for f in p["filas"]] == [
            ("Residuales", "L1", "W45", 1), ("Residuales", "L1", "W46", 2),
            ("Patologías", "L2", "W46", 4), ("Patologías", "L2", "W47", 4),
        ]
        assert p["descuadres"] == []

    def test_los_totales_del_pie_y_la_simbologia_no_son_clientes(self):
        p = leer_plan(_excel([("AGROVIC", "Residuales L1", [1, 1, 1], 3)]))
        assert {f["cliente_plan"] for f in p["filas"]} == {"AGROVIC"}

    def test_la_linea_pegada_al_cliente_se_separa(self):
        f = leer_plan(_excel([("AURORA AUSTRALIS L2", "Residuales ", [1, None, None], 1)]))["filas"][0]
        assert (f["cliente_plan"], f["linea"], f["servicio"]) == ("AURORA AUSTRALIS", "L2", "Residuales")

    def test_avisa_cuando_el_total_no_cuadra_con_las_semanas(self):
        p = leer_plan(_excel([("DOSAL", "Residuales ", [1, 2, 2], 4)]))  # el Excel dice 4, las semanas suman 5
        assert p["descuadres"] == [{"fila_excel": 6, "cliente_plan": "DOSAL", "servicio": "Residuales", "total_excel": 4, "suma_semanas": 5}]

    def test_semanas_iso_de_la_temporada(self):
        assert lunes_de_semana("W45") == date(2026, 11, 2)
        assert lunes_de_semana("W01") == date(2027, 1, 4)

    def test_categorias_de_servicio(self):
        for texto, esperado in [
            ("Residuales L1", "Residuales"), ("Patologias L2", "Patologías"), ("Fitopatológico", "Patologías"),
            ("Microbilogicos", "Microbiología"), ("Microbiológicos", "Microbiología"), ("Registros", "Entrega de registros"), ("Fsma L1", "FSMA"),
        ]:
            assert categoria_servicio(texto) == esperado


CONTACTOS = [
    {"tipo": "resultado_interno", "cargo": "Comercial", "tipo_copia": "cc", "email": "jlabarca@agrofresh.com", "nombre": "JC Labarca",
     "sold_to": "DOLE CHILE S.A.", "ship_to": "DOLE LONTUE", "activo": True},
    {"tipo": "resultado_interno", "cargo": "Comercial", "tipo_copia": "cc", "email": "jlabarca@agrofresh.com", "nombre": "JC Labarca",
     "sold_to": "EXPORTADORA UNIFRUTTI TRADERS", "ship_to": "UNIFRUTTI LINARES", "activo": True},
    {"tipo": "resultado_interno", "cargo": "Comercial", "tipo_copia": "cc", "email": "rtraverso@agrofresh.com", "nombre": "Rodrigo Traverso",
     "sold_to": "EXPORTADORA PRIZE S.A", "ship_to": "EXPORTADORA PRIZE LTDA.", "activo": True},
    {"tipo": "resultado_interno", "cargo": "Técnico", "tipo_copia": "bcc", "email": "tecnico@agrofresh.com", "nombre": "Técnico",
     "sold_to": "DOLE CHILE S.A.", "ship_to": "DOLE LONTUE", "activo": True},
]


class TestCruce:
    def test_identifica_al_comercial_por_su_apellido(self):
        internos = [c for c in CONTACTOS if c["cargo"] == "Comercial"]
        assert correos_de_comercial("JC Labarca", internos) == ["jlabarca@agrofresh.com"]
        assert correos_de_comercial("R. TRAVERSO", internos) == ["rtraverso@agrofresh.com"]
        assert correos_de_comercial("Alguien Distinto", internos) == []

    def test_las_plantas_del_comercial_salen_de_la_lista(self):
        estado = estado_desde_contactos(CONTACTOS)
        assert plantas_de(estado, ["jlabarca@agrofresh.com"]) == [
            ("DOLE CHILE S.A.", "DOLE LONTUE"), ("EXPORTADORA UNIFRUTTI TRADERS", "UNIFRUTTI LINARES"),
        ]
        assert plantas_de(estado, ["rtraverso@agrofresh.com"]) == [("EXPORTADORA PRIZE S.A", "EXPORTADORA PRIZE LTDA.")]

    def test_igual_cuando_el_nombre_es_el_del_ship_to(self):
        r = proponer("DOLE LONTUE", [("DOLE CHILE S.A.", "DOLE LONTUE"), ("DOLE CHILE S.A.", "DOLE CODEGUA")])
        assert r["estado"] == "Igual" and r["propuesta"] == ("DOLE CHILE S.A.", "DOLE LONTUE")

    def test_probable_cuando_una_sola_planta_contiene_las_palabras(self):
        r = proponer("PRIZE", [("EXPORTADORA PRIZE S.A", "EXPORTADORA PRIZE LTDA."), ("DOLE CHILE S.A.", "DOLE LONTUE")])
        assert r["estado"] == "Probable" and r["propuesta"][1] == "EXPORTADORA PRIZE LTDA."

    def test_varias_cuando_no_se_puede_elegir(self):
        r = proponer("UNIFRUTTI", [("EXPORTADORA UNIFRUTTI TRADERS", "UNIFRUTTI LINARES"), ("EXPORTADORA UNIFRUTTI TRADERS", "UNIFRUTTI LINDEROS")])
        assert r["estado"] == "Varias" and r["propuesta"] is None and len(r["candidatas"]) == 2

    def test_revisar_cuando_solo_comparten_la_marca(self):
        r = proponer("UNIFRUTTI TENO", [("EXPORTADORA UNIFRUTTI TRADERS", "UNIFRUTTI LINARES")])
        assert r["estado"] == "Revisar" and r["propuesta"] is None

    def test_sin_candidato(self):
        assert proponer("ZUR GRUP", [("DOLE CHILE S.A.", "DOLE LONTUE")])["estado"] == "Sin candidato"

    def test_no_adivina_solo_por_el_sold_to(self):
        # «DOLE SAN FELIPE» comparte «DOLE» con el Sold To, pero ninguna planta es San Felipe
        r = proponer("DOLE SAN FELIPE", [("DOLE CHILE S.A.", "DOLE LONTUE")])
        assert r["propuesta"] is None

    def test_una_coincidencia_solo_del_sold_to_se_revisa_no_se_da_por_buena(self):
        r = proponer("GREENEX", [("GREENEX SPA", "")])
        assert r["estado"] == "Revisar" and r["propuesta"] is None


class TestUnificado:
    def _unificado(self) -> io.BytesIO:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Entregables"
        ws.append(["Comercial", "Cliente (como en el plan)", "Línea", "Servicio", "Servicio (original)", "Semana", "Desde", "Hasta", "Cantidad", "Archivo", "Fila en el Excel"])
        ws.append(["JC Labarca", "DOLE LONTUE", "L1", "Residuales", "Residuales L1", "W45", "2026-11-02", "2026-11-08", 2, "Entregables_JCL.xlsx", 6])
        ws.append(["JC Labarca", "DOLE LONTUE", "L1", "Residuales", "Residuales L1", "W46", "2026-11-09", "2026-11-15", 3, "Entregables_JCL.xlsx", 6])
        ws.append(["R. TRAVERSO", "PRIZE", "", "Residuales", "Residuales", "W45", "2026-11-02", "2026-11-08", 1, "Entregables_RT.xlsx", 6])
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf

    def test_el_excel_unificado_se_separa_por_comercial(self):
        planes = leer_planes(self._unificado())
        assert [p["comercial"] for p in planes] == ["JC Labarca", "R. TRAVERSO"]
        assert [(f["cliente_plan"], f["semana"], f["cantidad"]) for f in planes[0]["filas"]] == [("DOLE LONTUE", "W45", 2), ("DOLE LONTUE", "W46", 3)]
        assert planes[0]["semanas"] == ["W45", "W46"]

    def test_un_plan_individual_sigue_leyendose_como_antes(self):
        planes = leer_planes(_excel([("DOLE LONTUE", "Residuales L1", [1, 2, None], 3)]))
        assert len(planes) == 1 and planes[0]["comercial"] == "JC Labarca" and len(planes[0]["filas"]) == 2
